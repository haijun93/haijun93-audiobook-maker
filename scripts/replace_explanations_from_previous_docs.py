#!/usr/bin/env python3
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
BASE_DIR = Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/260416")
PREV_DIR = BASE_DIR / "이전"
BACKUP_DIR = BASE_DIR / "backup"


CONFIGS = [
    {
        "name": "social",
        "docx": BASE_DIR / "[doc] social_insurance_law_OX_integrated_2026_v1.docx",
        "reference": PREV_DIR / "social-insurance-law-260405_(1).docx",
        "stopwords": {
            "사회보험법", "사회보장기본법", "고용보험법", "산업재해보상보험법", "국민연금법", "국민건강보험법",
            "징수법", "문제", "지문", "설명", "기준", "법률", "조문", "국민", "국가", "지방자치단체",
        },
    },
    {
        "name": "labor",
        "docx": BASE_DIR / "[doc] labor_law_OX_integrated_2026_v1.docx",
        "reference": PREV_DIR / "#labor_law_2.docx",
        "stopwords": {
            "노동법", "헌법", "근로기준법", "노동조합", "노동위원회", "문제", "지문", "설명", "기준", "조문",
        },
    },
    {
        "name": "civil",
        "docx": BASE_DIR / "[doc] civil_law_OX_integrated_2026_v1.docx",
        "reference": PREV_DIR / "civil_law_260405_1.docx",
        "stopwords": {
            "민법", "문제", "지문", "설명", "기준", "판례", "조문", "당사자", "효과", "요건",
        },
    },
    {
        "name": "management",
        "docx": BASE_DIR / "[doc] management_OX_integrated_2026_v1.docx",
        "reference": PREV_DIR / "bs admin_260405_1.docx",
        "stopwords": {
            "경영학", "개론", "문제", "지문", "설명", "기준", "개념", "정의", "내용", "전략", "관리",
        },
    },
]


@dataclass
class Candidate:
    text: str
    keywords: set[str]
    context: str


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def compact(text: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣#]+", "", text)


def extract_basis_focus(basis: str) -> str:
    basis = normalize(basis)
    match = re.search(r"\(([^()]+)\)", basis)
    if match:
        return match.group(1).strip()
    match = re.search(r"(제\d+조(?:의\d+)?(?:\s*제\d+항)?)", basis)
    if match:
        return match.group(1).strip()
    return ""


def keyword_set(text: str, stopwords: set[str]) -> set[str]:
    words: set[str] = set()
    common_stopwords = {
        "원칙", "경우", "관계", "정하는", "한다", "아니한다", "있다", "없다", "기준", "구조", "해당", "다음",
        "포함", "구분", "관련", "따른", "의한", "대한", "통한", "각각",
    }
    endings = [
        "하여야한다", "해야한다", "하지않는다", "하지아니한다", "되어야한다", "되지아니한다",
        "하여야", "해야", "한다", "된다", "되는", "되는지", "되며", "하는", "하여", "되고", "이다",
        "이며", "으로", "에는", "에서", "에게", "까지", "부터", "보다", "처럼", "이면", "만", "도",
        "은", "는", "이", "가", "을", "를", "의", "와", "과", "로", "에",
    ]
    for word in re.findall(r"[A-Za-z]+|[가-힣]{2,}|\d+", normalize(text)):
        if len(word) <= 1 or word in stopwords or word in common_stopwords:
            continue
        words.add(word)
        if re.fullmatch(r"[가-힣]+", word):
            stem = word
            for ending in endings:
                if stem.endswith(ending) and len(stem) - len(ending) >= 2:
                    stem = stem[: -len(ending)]
                    break
            if stem not in stopwords and stem not in common_stopwords and len(stem) >= 2:
                words.add(stem)
    aliases: dict[str, tuple[str, ...]] = {
        "대리": ("대신", "대신신청"),
        "대리하여": ("대신", "대신신청"),
        "대신": ("대리", "대신신청"),
        "직권": ("대신신청",),
        "신청일": ("이송",),
        "압류": ("양도", "담보"),
        "담보": ("압류", "양도"),
        "양도": ("압류", "담보"),
        "상실": ("소멸",),
        "소멸": ("상실",),
        "정지": ("소멸", "상실"),
        "재심사": ("재심",),
    }
    expanded = set(words)
    for word in list(words):
        for key, alias_words in aliases.items():
            if key in word:
                expanded.update(alias_words)
    return expanded


def iter_doc_lines(doc: Document) -> list[str]:
    lines: list[str] = []
    for p in doc.paragraphs:
        raw = p.text.replace("\xa0", " ")
        for part in re.split(r"[\n\r\v]+", raw):
            text = normalize(part)
            if text:
                lines.append(text)
    return lines


def is_heading_line(text: str) -> bool:
    t = normalize(text)
    if not t:
        return False
    if t.startswith("•") or t.startswith("판례 원문") or t.startswith("정리:") or t.startswith("[포인트]"):
        return False
    if t.startswith("[") and t.endswith("]"):
        return False
    if len(t) > 32:
        return False
    if t.endswith("다.") or t.endswith("다") and len(t) > 20:
        return False
    banned_starts = (
        "공인노무사", "학습 사용법", "빠른 암기 네비게이션", "전체 구조", "리디자인 포인트",
        "출력·제본", "대표 두문자 네비게이션", "이번 장 5초 회상", "PART ", "Ⅰ.", "Ⅱ.", "Ⅲ.", "Ⅳ.", "Ⅴ.", "Ⅵ.",
    )
    if t.startswith(banned_starts):
        return False
    if t in {"판례 보강", "관리사상 흐름", "동기이론 빠른 구분", "리더십 이론 빠른 구분"}:
        return False
    return True


def is_useful_line(text: str) -> bool:
    t = normalize(text)
    if len(t) < 6:
        return False
    banned_starts = (
        "공인노무사", "학습 사용법", "빠른 암기 네비게이션", "전체 구조", "리디자인 포인트",
        "출력·제본", "대표 두문자 네비게이션", "이번 장 5초 회상", "PART ", "Ⅰ.", "Ⅱ.", "Ⅲ.", "Ⅳ.", "Ⅴ.", "Ⅵ.",
    )
    if t.startswith(banned_starts):
        return False
    if t in {"[비교]", "[숫자]", "판례 보강", "관리사상 흐름", "동기이론 빠른 구분", "리더십 이론 빠른 구분"}:
        return False
    if t.startswith("•") or t.startswith("판례 원문") or t.startswith("정리:") or t.startswith("[포인트]"):
        return True
    if "=" in t or ":" in t[:18]:
        return True
    if len(t) >= 20 and t.endswith("다."):
        return True
    return False


def load_candidates(reference_path: Path, stopwords: set[str]) -> list[Candidate]:
    doc = Document(reference_path)
    lines = iter_doc_lines(doc)

    candidates: list[Candidate] = []
    seen: set[str] = set()
    heading_chain: list[str] = []

    def add_candidate(text: str, context: str) -> None:
        text = normalize(text)
        if not text or text in seen:
            return
        keys = keyword_set(text + " " + context, stopwords)
        if not keys and len(text) < 12:
            return
        candidates.append(Candidate(text=text, keywords=keys, context=context))
        seen.add(text)

    for line in lines:
        if is_heading_line(line):
            if not heading_chain or heading_chain[-1] != line:
                heading_chain.append(line)
                heading_chain = heading_chain[-2:]
            continue
        if is_useful_line(line):
            add_candidate(line, " ".join(heading_chain))
    return candidates


def build_index(candidates: list[Candidate]) -> dict[str, list[int]]:
    index: dict[str, list[int]] = {}
    for idx, cand in enumerate(candidates):
        for key in cand.keywords:
            index.setdefault(key, []).append(idx)
    return index


def score_candidate(question_text: str, basis: str, candidate: Candidate, stopwords: set[str]) -> tuple[float, int, int]:
    focus = extract_basis_focus(basis)
    qkeys = keyword_set(question_text + " " + basis, stopwords)
    query_text = normalize(question_text + " " + basis)
    overlap = len(qkeys & candidate.keywords)
    score = float(overlap)
    focus_hits = 0
    if focus:
        focus_words = keyword_set(focus, stopwords)
        focus_hits = len(focus_words & candidate.keywords)
        score += 2.0 * focus_hits
        if focus in candidate.text:
            score += 2.5
        if focus and focus in candidate.context:
            score += 1.5
    for word in sorted(qkeys, key=len, reverse=True)[:5]:
        if len(word) >= 2 and word in candidate.text:
            score += 0.7
        elif len(word) >= 2 and word in candidate.context:
            score += 0.25
    rare_hits = sum(1 for word in qkeys if len(word) >= 3 and word in candidate.text)
    score += rare_hits * 0.35
    if any(token in query_text for token in ("대리", "직권")) and "대신" in candidate.text:
        score += 2.0
    if any(token in query_text for token in ("이송", "신청일")) and any(token in candidate.text for token in ("이송", "신청일")):
        score += 2.0
    if any(token in query_text for token in ("압류", "담보", "양도")) and any(token in candidate.text for token in ("압류", "담보", "양도")):
        score += 1.0
    return score, overlap, focus_hits


def rank_references(
    question_text: str,
    basis: str,
    candidates: list[Candidate],
    keyword_index: dict[str, list[int]],
    stopwords: set[str],
) -> list[tuple[float, int, int, str]]:
    qkeys = keyword_set(question_text + " " + basis, stopwords)
    candidate_ids: set[int] = set()
    for key in qkeys:
        candidate_ids.update(keyword_index.get(key, []))
    if not candidate_ids:
        return []

    ranked: list[tuple[float, int, int, str]] = []
    for idx in candidate_ids:
        cand = candidates[idx]
        score, overlap, focus_hits = score_candidate(question_text, basis, cand, stopwords)
        ranked.append((score, overlap, focus_hits, cand.text))
    ranked.sort(key=lambda x: x[0], reverse=True)
    return ranked


def choose_reference(
    question_text: str,
    basis: str,
    candidates: list[Candidate],
    keyword_index: dict[str, list[int]],
    stopwords: set[str],
    avoid_text: str = "",
) -> str:
    ranked = rank_references(question_text, basis, candidates, keyword_index, stopwords)
    if not ranked:
        return ""
    best_score, best_overlap, best_focus_hits, best_text = ranked[0]
    if best_score < 2.2:
        return ""
    focus = extract_basis_focus(basis)
    if focus and best_focus_hits == 0 and best_overlap < 1:
        return ""
    if best_overlap < 1:
        return ""
    if avoid_text:
        for score, overlap, focus_hits, text in ranked[1:]:
            if text == avoid_text:
                continue
            if score < best_score - 0.75:
                break
            if focus and focus_hits == 0 and overlap < 1:
                continue
            if overlap < 1:
                continue
            return text.strip()
    return best_text.strip()


def backup_file(path: Path) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"{path.stem}_backup_before_reference_only_{ts}{path.suffix}"
    backup.write_bytes(path.read_bytes())
    return backup


def replace_explanations(docx_path: Path, reference_path: Path, stopwords: set[str]) -> tuple[Path, int]:
    candidates = load_candidates(reference_path, stopwords)
    keyword_index = build_index(candidates)
    backup = backup_file(docx_path)
    doc = Document(docx_path)

    current_question = ""
    current_basis = ""
    previous_ref_text = ""
    updated = 0

    for p in doc.paragraphs:
        text = normalize(p.text)
        if re.match(r"^\d+\.\s", text):
            current_question = text.split(". ", 1)[1].strip()
        elif text.startswith("근거:"):
            current_basis = text.split(":", 1)[1].strip()
        elif text.startswith("해설:"):
            ref_text = choose_reference(
                current_question,
                current_basis,
                candidates,
                keyword_index,
                stopwords,
                avoid_text=previous_ref_text,
            )
            p.text = f"해설: {ref_text}" if ref_text else "해설:"
            previous_ref_text = ref_text
            updated += 1

    doc.save(docx_path)
    return backup, updated


def main() -> None:
    for cfg in CONFIGS:
        backup, updated = replace_explanations(cfg["docx"], cfg["reference"], cfg["stopwords"])
        print(f"{cfg['name']}: backup={backup} updated={updated}")


if __name__ == "__main__":
    main()
