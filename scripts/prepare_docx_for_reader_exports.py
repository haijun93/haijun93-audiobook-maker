#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
import runpy
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]


REMOVE_EXACT_LINES = {
    "정답은 9개씩 묶어서 확인하면 오답 패턴이 더 잘 보인다.",
    "OX 정답표",
    "최종 정오표",
    "정답표 · 최종 정오표 정답은 9개씩 묶어서 확인하면 오답 패턴이 더 잘 보인다.",
    "OX 정답표 구간 정답 최종 정오표 항목 최종 확인",
}


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="DOCX를 학습용으로 정리하고, PDF는 Scribe용 / EPUB은 6인치 Kindle용으로 내보냅니다."
    )
    parser.add_argument("docx_files", nargs="+", type=Path, help="정리할 DOCX 파일")
    return parser.parse_args()


def set_run_font(run, size_pt: float, *, bold: bool | None = None) -> None:
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def normalize_style_fonts(doc: Document) -> None:
    for name, size, bold in [
        ("Normal", 10.5, None),
        ("Title", 20, True),
        ("Heading 1", 13.5, True),
        ("Heading 2", 12.5, True),
        ("Heading 3", 11.5, True),
    ]:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold


def paragraph_role(text: str, nonempty_index: int) -> tuple[float, bool | None, int | None]:
    stripped = text.strip()
    if nonempty_index == 0:
        return 22.0, True, WD_ALIGN_PARAGRAPH.CENTER
    if nonempty_index in {1, 2} and len(stripped) <= 55 and not re.match(r"^\d+\.", stripped):
        return 11.5, True, WD_ALIGN_PARAGRAPH.CENTER
    if re.match(r"^\d+\.\s", stripped):
        return 13.0, True, None
    if stripped.endswith(("체크표", "한눈표", "핵심", "정리", "총정리")) and len(stripped) <= 40:
        return 12.0, True, None
    if stripped.startswith(("•", "-", "▣", "☞")):
        return 10.5, None, None
    return 10.5, None, None


def is_revision_points_table(table) -> bool:
    if not table.rows:
        return False
    header_cells = [cell.text.strip() for cell in table.rows[0].cells]
    needed = {"시점", "주제", "핵심 변화", "암기 포인트"}
    return needed.issubset(set(header_cells))


def is_numbers_overview_table(table) -> bool:
    if not table.rows:
        return False
    header_cells = [cell.text.strip() for cell in table.rows[0].cells]
    needed = {"숫자/날짜", "키워드", "암기 포인트"}
    return needed.issubset(set(header_cells))


def insert_paragraph_after(anchor: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    new_para = Paragraph(new_p, anchor._parent)
    new_para.add_run(text)
    return new_para


def format_inserted_summary_paragraph(para: Paragraph, *, size_pt: float = 10.2) -> None:
    pf = para.paragraph_format
    pf.line_spacing = 1.08
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    for run in para.runs:
        set_run_font(run, size_pt)


def replace_target_table_with_text(doc: Document, heading_text: str, match_fn, row_builder) -> bool:
    body = doc._body._body
    children = list(body.iterchildren())
    heading_para = None
    table_obj = None

    for idx, child in enumerate(children):
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, doc._body)
            if para.text.strip() == heading_text:
                heading_para = para
                # find next table sibling
                for next_child in children[idx + 1 :]:
                    next_tag = next_child.tag.split("}")[-1]
                    if next_tag == "tbl":
                        candidate = Table(next_child, doc._body)
                        if match_fn(candidate):
                            table_obj = candidate
                        break
                    if next_tag == "p" and Paragraph(next_child, doc._body).text.strip():
                        break
                break

    if heading_para is None or table_obj is None:
        return False

    current = heading_para
    for row in table_obj.rows[1:]:
        text = row_builder([cell.text.strip().replace("\n", " / ") for cell in row.cells])
        para = insert_paragraph_after(current, text)
        format_inserted_summary_paragraph(para)
        current = para

    tbl_el = table_obj._element
    tbl_el.getparent().remove(tbl_el)
    return True


def clean_docx_in_place(docx_path: Path) -> Path:
    backup = docx_path.with_name(f"{docx_path.stem}_backup_before_reader_prep_{timestamp()}{docx_path.suffix}")
    shutil.copy2(docx_path, backup)

    doc = Document(str(docx_path))
    normalize_style_fonts(doc)

    nonempty_seen = 0
    removed_count = 0
    for para in list(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        if text in REMOVE_EXACT_LINES:
            p = para._element
            p.getparent().remove(p)
            removed_count += 1
            continue

        size_pt, bold, align = paragraph_role(text, nonempty_seen)
        nonempty_seen += 1
        if align is not None:
            para.alignment = align
        elif para.alignment == WD_ALIGN_PARAGRAPH.CENTER and nonempty_seen > 3:
            para.alignment = None

        pf = para.paragraph_format
        pf.line_spacing = 1.12
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)

        for run in para.runs:
            set_run_font(run, size_pt, bold=bold)

    if "civil law_260420" in docx_path.name.lower():
        replace_target_table_with_text(
            doc,
            "개정 포인트 체크표",
            is_revision_points_table,
            lambda cells: f"• {cells[0]} | {cells[1]}: {cells[2]} / 암기 포인트: {cells[3]}",
        )
        replace_target_table_with_text(
            doc,
            "숫자·비율·기간 한눈표",
            is_numbers_overview_table,
            lambda cells: f"• {cells[0]} | {cells[1]}: {cells[2]}",
        )

    for table in doc.tables:
        small_table = is_revision_points_table(table)
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    pf = para.paragraph_format
                    pf.line_spacing = 0.98 if small_table else 1.05
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    for run in para.runs:
                        set_run_font(run, 8.4 if small_table else 9.2)

    doc.save(str(docx_path))
    print(f"cleaned_docx={docx_path}")
    print(f"backup_docx={backup}")
    print(f"removed_lines={removed_count}")
    return backup


def export_outputs(docx_path: Path) -> tuple[Path, Path]:
    optimize_helpers = runpy.run_path(str(ROOT / "scripts" / "optimize_pdfs_for_kindle_scribe.py"))
    export_helpers = runpy.run_path(str(ROOT / "scripts" / "export_docx_to_pdf_epub.py"))

    optimize_docx_for_kindle = optimize_helpers["optimize_docx_for_kindle"]
    export_docx_with_pages = export_helpers["export_docx_with_pages"]
    render_pdf_first_page_to_png_bytes = export_helpers["render_pdf_first_page_to_png_bytes"]
    ensure_epub_cover_from_png = export_helpers["ensure_epub_cover_from_png"]

    pdf_path = docx_path.with_name(f"[pdf] {docx_path.stem}.pdf")
    epub_path = docx_path.with_name(f"[epub] {docx_path.stem}.epub")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        pdf_docx = tmp / f"scribe_{docx_path.name}"
        optimize_docx_for_kindle(docx_path, pdf_docx, profile="scribe")
        export_docx_with_pages(pdf_docx, pdf_path, tmp / "unused.epub")

        epub_docx = tmp / f"kindle6_{docx_path.name}"
        optimize_docx_for_kindle(docx_path, epub_docx, profile="kindle6")
        temp_pdf_for_cover = tmp / "kindle6_cover.pdf"
        export_docx_with_pages(epub_docx, temp_pdf_for_cover, epub_path)
        cover_png = render_pdf_first_page_to_png_bytes(temp_pdf_for_cover, dpi=144)
        ensure_epub_cover_from_png(epub_path, cover_png)

    print(f"pdf={pdf_path}")
    print(f"epub={epub_path}")
    return pdf_path, epub_path


def main() -> int:
    args = parse_args()
    files = [p.expanduser().resolve() for p in args.docx_files]
    for path in files:
        if not path.exists():
            raise SystemExit(f"입력 DOCX를 찾지 못했습니다: {path}")
    for path in files:
        clean_docx_in_place(path)
        export_outputs(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
