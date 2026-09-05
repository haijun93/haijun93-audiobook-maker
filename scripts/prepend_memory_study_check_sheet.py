#!/usr/bin/env python3

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt


BASE_DIR = Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416")
BACKUP_DIR = BASE_DIR / "backup"
SUBJECT_DOCS = [
    BASE_DIR / "[doc] labor_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] social_insurance_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] civil_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] management_OX_integrated_2026_v1.docx",
]
MASTER_GUIDE = BASE_DIR / "[doc] 35day_first_exam_master_guide_2026.docx"


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def set_run_font(run, name: str = "Malgun Gothic", size: float = 10.5, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def backup_file(path: Path) -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup = BACKUP_DIR / f"{path.stem}_backup_before_memory_sheet_{timestamp()}{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def pop_body_elements(doc: Document):
    body = doc._body._element
    elements = []
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        elements.append(child)
        body.remove(child)
    return elements


def append_body_elements(doc: Document, elements) -> None:
    body = doc._body._element
    sect = None
    children = list(body)
    for child in children:
        if child.tag.endswith("sectPr"):
            sect = child
            break
    insert_index = children.index(sect) if sect is not None else len(children)
    for element in elements:
        body.insert(insert_index, element)
        insert_index += 1


def add_paragraph(doc: Document, text: str, *, style: str | None = None, bold: bool = False, align=None) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, bold=bold, size=12 if style == "Heading 1" else 10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_subject_sheet(doc: Document) -> None:
    add_paragraph(doc, "과학적 기억암기 기반 학습체크 기록지", style="Heading 1")
    add_bullets(
        doc,
        [
            "이 체크지는 간격 반복, 인출 연습, 교차 회독, 오답 재부호화 원리에 맞춰 설계했다.",
            "학습 직후보다 D+1, D+3, D+7 복습에서 기억이 가장 많이 남는다. 체크는 '완료 여부'보다 '책을 덮고 말로 설명했는가'를 기준으로 한다.",
            "한 번 읽고 넘어가지 말고, 각 단원마다 아래 표를 따라 최소 6회 이상 다시 보아야 점수로 연결된다.",
        ],
    )

    add_paragraph(doc, "1. 회독 루틴", style="Heading 2")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["회차", "권장 시점", "무엇을 할지", "체크", "완료 날짜"]
    for i, header in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(header)
        set_run_font(run, bold=True)
    rows = [
        ["1차", "D0", "개념·표·해설 읽기 후 바로 구조 잡기", "□", ""],
        ["즉시 회상", "학습 직후 10분", "책을 덮고 목차·두문자·숫자·절차를 말로 복원", "□", ""],
        ["2차", "D+1", "전날 범위 OX 재확인, 틀린 문항만 표시", "□", ""],
        ["3차", "D+3", "해설 안 보고 정답 이유를 말로 설명", "□", ""],
        ["4차", "D+7", "숫자·위원회·절차만 압축 회독", "□", ""],
        ["5차", "D+14", "시간 제한으로 다시 풀고 오답만 재정리", "□", ""],
        ["6차", "D+30", "최종 고정 회독", "□", ""],
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            run = cells[idx].paragraphs[0].add_run(value)
            set_run_font(run)

    add_paragraph(doc, "2. 단원별 체크 기록", style="Heading 2")
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["단원/범위", "D0", "즉시회상", "D+1", "D+3", "D+7", "D+14", "메모"]
    for i, header in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(header)
        set_run_font(run, bold=True)
    for _ in range(12):
        cells = table.add_row().cells
        for idx in range(8):
            value = "□" if 1 <= idx <= 6 else ""
            run = cells[idx].paragraphs[0].add_run(value)
            set_run_font(run)

    add_paragraph(doc, "3. 사용 원칙", style="Heading 2")
    add_bullets(
        doc,
        [
            "체크는 읽었는지보다 '말로 설명했는지'를 기준으로 한다.",
            "같은 날 너무 많이 밀리지 않게, 새 범위 1개를 보면 반드시 직전 범위 D+1 또는 D+3 복습 1개를 같이 수행한다.",
            "메모 칸에는 긴 오답노트 대신 '#두문자', 숫자, 위원회 이름, 틀린 이유 한 줄만 적는다.",
        ],
    )

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_master_sheet(doc: Document) -> None:
    add_paragraph(doc, "과학적 기억암기 기반 35일 체크 기록지", style="Heading 1")
    add_bullets(
        doc,
        [
            "원리는 간격 반복(spaced repetition), 인출 연습(retrieval practice), 교차 학습(interleaving), 오답 재부호화(error encoding)다.",
            "하루 학습은 '새 범위'보다 '재노출 주기 유지'가 더 중요하다. 오늘 배운 것은 오늘 10분 후, 내일, 3일 후, 7일 후 다시 보아야 남는다.",
            "각 날짜에는 새 학습과 복습이 함께 있어야 한다. 새로 공부한 범위만 체크하고 복습 칸을 비워 두면 효율이 급격히 떨어진다.",
        ],
    )

    add_paragraph(doc, "1. 35일 일정 체크표", style="Heading 2")
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["Day", "날짜", "핵심 과목", "새 학습 범위", "즉시회상", "D+1/D+3", "D+7/D+14", "메모"]
    for i, header in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(header)
        set_run_font(run, bold=True)
    for day in range(1, 36):
        cells = table.add_row().cells
        values = [str(day), "", "", "", "□", "□", "□", ""]
        for idx, value in enumerate(values):
            run = cells[idx].paragraphs[0].add_run(value)
            set_run_font(run)

    add_paragraph(doc, "2. 체크 기준", style="Heading 2")
    add_bullets(
        doc,
        [
            "즉시회상: 책을 덮고 3분 안에 목차·핵심요건·#두문자를 말할 수 있으면 체크한다.",
            "D+1/D+3: 해설을 다시 읽지 말고 정답 이유를 먼저 떠올린 뒤 확인한다.",
            "D+7/D+14: 숫자, 위원회, 신청·재심 절차처럼 헷갈리는 비교 포인트를 표 중심으로 회독한다.",
            "메모는 길게 쓰지 말고 '틀린 이유 1줄 + #두문자 1개'만 남긴다.",
        ],
    )

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def process_doc(path: Path, is_master: bool) -> Path:
    backup = backup_file(path)
    doc = Document(path)
    existing = pop_body_elements(doc)
    if is_master:
        add_master_sheet(doc)
    else:
        add_subject_sheet(doc)
    append_body_elements(doc, existing)
    doc.save(path)
    return backup


def main() -> int:
    for path in SUBJECT_DOCS + [MASTER_GUIDE]:
        if not path.exists():
            raise SystemExit(f"파일을 찾지 못했습니다: {path}")

    for path in SUBJECT_DOCS:
        backup = process_doc(path, is_master=False)
        print(f"{path}\nbackup={backup}")
    backup = process_doc(MASTER_GUIDE, is_master=True)
    print(f"{MASTER_GUIDE}\nbackup={backup}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
