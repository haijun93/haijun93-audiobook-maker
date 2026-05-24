#!/usr/bin/env python3
from __future__ import annotations

import re
import runpy
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document


BASE_DIR = Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/260416")
BACKUP_DIR = BASE_DIR / "backup"

CONFIGS = [
    {
        "name": "labor",
        "docx": BASE_DIR / "[doc] labor_law_OX_integrated_2026_v1.docx",
        "pdf": BASE_DIR / "[pdf] labor_law_OX_integrated_2026_v1.pdf",
        "source": BACKUP_DIR / "labor_law_OX_integrated_2026_v1_backup_before_enhance_20260416_222931.docx",
    },
    {
        "name": "civil",
        "docx": BASE_DIR / "[doc] civil_law_OX_integrated_2026_v1.docx",
        "pdf": BASE_DIR / "[pdf] civil_law_OX_integrated_2026_v1.pdf",
        "source": BACKUP_DIR / "civil_law_OX_integrated_2026_v1_backup_before_enhance_20260416_222931.docx",
    },
    {
        "name": "management",
        "docx": BASE_DIR / "[doc] management_OX_integrated_2026_v1.docx",
        "pdf": BASE_DIR / "[pdf] management_OX_integrated_2026_v1.pdf",
        "source": BACKUP_DIR / "management_OX_integrated_2026_v1_backup_before_enhance_20260416_222931.docx",
    },
]


TRIGGER_PHRASES = [
    "아니다",
    "할 수 없다",
    "할 수 있다",
    "무효이다",
    "취소할 수 있다",
    "철회할 수 있다",
    "반드시",
    "언제나",
    "당연히",
    "직권으로만",
    "전액",
    "일체",
]


def ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def strip_bullet(text: str) -> str:
    text = normalize(text)
    text = re.sub(r"^[•·]\s*", "", text)
    return text


def backup_file(path: Path, suffix: str) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    backup = BACKUP_DIR / f"{path.stem}_{suffix}_{ts()}{path.suffix}"
    backup.write_bytes(path.read_bytes())
    return backup


def build_direct_x_map(docx_path: Path) -> dict[int, str]:
    doc = Document(docx_path)
    direct_map: dict[int, str] = {}
    current_q: int | None = None
    current_answer = ""
    for p in doc.paragraphs:
        text = normalize(p.text)
        if re.match(r"^\d+\.\s", text):
            current_q = int(text.split(".", 1)[0])
            current_answer = ""
            continue
        if text.startswith(("정답:", "● 정답:")):
            current_answer = text.split(":", 1)[1].strip()
            continue
        if current_q is None or current_answer != "X" or not text.startswith("해설:"):
            continue
        body = text.split(":", 1)[1].strip()
        body = re.sub(r"^정답은 X입니다\.\s*", "", body)
        body = re.split(r"\s*(?:핵심:|기억:|두문자:|암기비법:|함께:)\s*", body, maxsplit=1)[0].strip()
        body = normalize(body).rstrip(".")
        if body:
            direct_map[current_q] = body
    return direct_map


def build_common_map(doc: Document) -> dict[str, str]:
    common_map: dict[str, str] = {}
    for p in doc.paragraphs:
        text = normalize(p.text)
        if text.startswith("공통해설(") and ":" in text:
            label, body = text.split(":", 1)
            common_map[label.strip()] = strip_bullet(body)
    return common_map


def resolve_current_explanation(text: str, common_map: dict[str, str]) -> str:
    text = normalize(text)
    if text.startswith("공통해설(") and ":" in text:
        return strip_bullet(text.split(":", 1)[1])
    match = re.match(r"^해설:\s*(공통해설\(문항 [^)]+\)) 참조$", text)
    if match:
        return common_map.get(match.group(1), "")
    if text.startswith("해설:"):
        return strip_bullet(text.split(":", 1)[1])
    return ""


def simplify_question(text: str) -> str:
    text = normalize(re.sub(r"^\d+\.\s*", "", text)).rstrip(".")
    prefixes = [
        "민법상 ",
        "근로기준법상 ",
        "노동조합 및 노동관계조정법상 ",
        "노동위원회법상 ",
        "국민건강보험법상 ",
        "경영학상 ",
        "경영학개론에서 ",
    ]
    for prefix in prefixes:
        if text.startswith(prefix):
            text = text[len(prefix):]
    return text


def extract_wrong_part(question_text: str, correction_text: str) -> str:
    question_text = simplify_question(question_text)
    correction_text = normalize(correction_text)

    for phrase in TRIGGER_PHRASES:
        if phrase in question_text:
            return f"'{phrase}'라고 한 부분이 틀렸다."

    token_pairs = [
        ("고용노동부장관", "중앙노동위원회"),
        ("중앙노동위원회", "지방노동위원회"),
        ("지방노동위원회", "중앙노동위원회"),
        ("무효", "취소"),
        ("취소", "무효"),
        ("철회", "취소"),
        ("표현대리", "무권대리"),
        ("공단", "심사평가원"),
    ]
    for wrong, right in token_pairs:
        if wrong in question_text and right in correction_text:
            return f"'{wrong}'이라고 한 부분이 틀렸다."

    if "없다" in question_text and "있다" in correction_text:
        return "'없다'고 단정한 부분이 틀렸다."
    if "있다" in question_text and "없다" in correction_text:
        return "'있다'고 한 부분이 틀렸다."

    if len(question_text) <= 72:
        return f"'{question_text}'라고 한 부분이 틀렸다."
    return "지문의 주체·요건·효과를 잘못 적은 부분이 틀렸다."


def format_explanation(question_text: str, correction_text: str) -> str:
    correction_text = strip_bullet(correction_text).rstrip(".")
    wrong_part = extract_wrong_part(question_text, correction_text)
    return f"해설: 틀린 부분: {wrong_part} 정정: {correction_text}."


def process_one(cfg: dict) -> tuple[Path, Path | None, int]:
    docx_path: Path = cfg["docx"]
    pdf_path: Path = cfg["pdf"]
    source_path: Path = cfg["source"]

    backup_doc = backup_file(docx_path, "backup_before_fix_x_explanations")
    backup_pdf = backup_file(pdf_path, "backup_before_fix_x_explanations") if pdf_path.exists() else None

    direct_x_map = build_direct_x_map(source_path)

    doc = Document(docx_path)
    common_map = build_common_map(doc)

    current_q: int | None = None
    current_answer = ""
    current_question_text = ""
    updated = 0

    for p in doc.paragraphs:
        text = normalize(p.text)
        if not text:
            continue
        if re.match(r"^\d+\.\s", text):
            current_q = int(text.split(".", 1)[0])
            current_question_text = text
            current_answer = ""
            continue
        if text.startswith(("정답:", "● 정답:")):
            current_answer = text.split(":", 1)[1].strip()
            continue
        if current_q is None or current_answer != "X":
            continue
        if not (text.startswith("해설:") or text.startswith("공통해설(")):
            continue

        correction_text = direct_x_map.get(current_q, "") or resolve_current_explanation(text, common_map)
        if not correction_text:
            correction_text = "지문의 문언이 정확한 기준과 다르므로 주체·요건·효과를 맞는 문장으로 고쳐야 한다"

        p.text = format_explanation(current_question_text, correction_text)
        updated += 1

    doc.save(docx_path)

    pdf_helpers = runpy.run_path("scripts/optimize_pdfs_for_kindle_scribe.py")
    optimize_docx_for_kindle = pdf_helpers["optimize_docx_for_kindle"]
    export_pdf_with_pages = pdf_helpers["export_pdf_with_pages"]

    with TemporaryDirectory() as tmpdir:
        temp_docx = Path(tmpdir) / docx_path.name
        optimize_docx_for_kindle(docx_path, temp_docx)
        export_pdf_with_pages(temp_docx, pdf_path)

    print(f"{cfg['name']}: backup_doc={backup_doc} backup_pdf={backup_pdf} updated={updated}")
    return backup_doc, backup_pdf, updated


def main() -> int:
    for cfg in CONFIGS:
        if not cfg["docx"].exists():
            raise SystemExit(f"문서를 찾지 못했습니다: {cfg['docx']}")
        if not cfg["source"].exists():
            raise SystemExit(f"참조 백업을 찾지 못했습니다: {cfg['source']}")
        process_one(cfg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
