#!/usr/bin/env python3
from __future__ import annotations

import re
import runpy
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document


BASE_DIR = Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416")
DOCX_PATH = BASE_DIR / "[doc] social_insurance_law_OX_integrated_2026_v1.docx"
PDF_PATH = BASE_DIR / "[pdf] social_insurance_law_OX_integrated_2026_v1.pdf"
BACKUP_DIR = BASE_DIR / "backup"
PREVIOUS_WITH_BASIS = BACKUP_DIR / "[doc] social_insurance_law_OX_integrated_2026_v1_before_remove_basis_20260417_224454.docx"
DIRECT_SOURCE_DOC = BACKUP_DIR / "social_insurance_law_OX_integrated_2026_v1_backup_before_enhance_20260416_222931.docx"


MANUAL_CORRECTIONS: dict[int, str] = {
    12: "사회보장정보시스템의 구축 및 운영 책임은 국무총리가 아니라 보건복지부장관에게 있다.",
    29: "국가 및 지방자치단체는 민간부문의 참여를 유도하기 위하여 정보 제공, 재정지원, 시설 지원 등을 할 수 있다.",
    38: "사회보장위원회 위원의 임기는 3년이 아니라 2년이며, 한 차례만 연임할 수 있다.",
    157: "장해급여 청구사유 발생 당시 대한민국 국민이 아닌 자로서 외국에 거주하는 근로자에게는 장해보상연금이 아니라 장해보상일시금을 지급한다.",
    211: "산업재해보상보험및예방기금은 근로복지공단 이사장이 아니라 고용노동부장관이 관리·운용하며, 수익성만이 아니라 안정성과 공공성을 함께 고려한다.",
    284: "국민연금사업은 고용노동부장관이 아니라 보건복지부장관이 맡아 주관한다.",
    345: "제4항에 따른 계약 조건 협의, 제5항에 따른 계약 체결 등에 필요한 사항은 대통령령이 아니라 보건복지부령으로 정한다.",
    346: "추천위원회는 대통령령이 아니라 보건복지부령으로 정하는 징수이사 후보 심사기준에 따라 심사하여야 하며, 징수이사 후보로 추천될 사람과 계약 조건을 협의하여야 한다.",
    347: "추천위원회는 주요 일간신문에 모집 공고를 하여야 하고, 적임자 조사나 전문단체 의뢰는 반드시 하여야 하는 것이 아니라 할 수 있다.",
    360: "요양급여비용의 심사와 요양급여의 적정성 평가는 국민건강보험공단이 아니라 건강보험심사평가원이 수행한다.",
    367: "국민건강보험공단의 주된 사무소 소재지는 법률이 아니라 정관으로 정한다.",
    368: "공단이 정관을 변경하려면 고용노동부장관이 아니라 보건복지부장관의 인가를 받아야 한다.",
    375: "공단의 예산안은 이사회 의결 후 고용노동부장관이 아니라 보건복지부장관의 승인을 받아야 하며, 예산 변경도 같다.",
    379: "건강보험사업은 고용노동부장관이 아니라 보건복지부장관이 맡아 주관한다.",
    394: "이사회의 의결 사항 및 운영 등에 필요한 사항은 보건복지부령이 아니라 대통령령으로 정한다.",
}


TRIGGER_PHRASES = [
    "의무적으로",
    "직권으로 신청하는 것은 어떠한 경우에도 허용되지 않는다",
    "직권으로 신청하는 것은",
    "국무총리에게 있다",
    "고용노동부장관이 맡아 주관한다",
    "보건복지부장관이 아니라",
    "압류하거나 담보로 제공할 수 있다",
    "제한, 정지, 양도, 압류가 어떠한 경우에도 일체 불가능",
    "구두 또는 서면으로 통지하여 포기할 수 있다",
    "서면으로 통지하여 포기하여야 한다",
    "재정지원을 할 수는 없다",
    "연임할 수 없다",
    "장해보상연금을 지급할 수 있다",
    "수익성을 가장 최우선적으로 고려하여야 한다",
    "대통령령으로 정한다",
    "법률로 명확히 정한다",
    "고용노동부장관의 인가를 받아야 한다",
    "건강보험공단이 수행한다",
]


def ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def backup_file(path: Path, suffix: str) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    backup = BACKUP_DIR / f"{path.stem}_{suffix}_{ts()}{path.suffix}"
    backup.write_bytes(path.read_bytes())
    return backup


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def strip_bullet(text: str) -> str:
    text = normalize(text)
    text = re.sub(r"^[•·]\s*", "", text)
    return text


def build_basis_map(docx_path: Path) -> dict[int, str]:
    doc = Document(docx_path)
    basis_map: dict[int, str] = {}
    current_q: int | None = None
    for p in doc.paragraphs:
        text = normalize(p.text)
        if re.match(r"^\d+\.\s", text):
            current_q = int(text.split(".", 1)[0])
        elif current_q is not None and text.startswith("근거:"):
            basis_map[current_q] = text.split(":", 1)[1].strip()
    return basis_map


def build_direct_x_map(docx_path: Path) -> dict[int, str]:
    doc = Document(docx_path)
    direct_map: dict[int, str] = {}
    current_q: int | None = None
    current_answer = ""
    for p in doc.paragraphs:
        text = normalize(p.text)
        if re.match(r"^\d+\.\s", text):
            current_q = int(text.split(".", 1)[0])
            current_answer = ""
            continue
        if text.startswith("정답:"):
            current_answer = text.split(":", 1)[1].strip()
            continue
        if current_q is None or current_answer != "X" or not text.startswith("해설:"):
            continue
        body = text.split(":", 1)[1].strip()
        body = re.sub(r"^정답은 X입니다\.\s*", "", body)
        body = re.split(r"\s*(?:핵심:|기억:|두문자:|암기비법:|함께:)\s*", body, maxsplit=1)[0].strip()
        body = normalize(body).rstrip(".")
        if body:
            direct_map[current_q] = body
    return direct_map


def build_common_map(doc: Document) -> dict[str, str]:
    common_map: dict[str, str] = {}
    for p in doc.paragraphs:
        text = normalize(p.text)
        if text.startswith("공통해설(") and ":" in text:
            label, body = text.split(":", 1)
            common_map[label.strip()] = strip_bullet(body)
    return common_map


def resolve_current_explanation(text: str, common_map: dict[str, str]) -> str:
    text = normalize(text)
    if text.startswith("공통해설(") and ":" in text:
        return strip_bullet(text.split(":", 1)[1])
    match = re.match(r"^해설:\s*(공통해설\(문항 [^)]+\)) 참조$", text)
    if match:
        return common_map.get(match.group(1), "")
    if text.startswith("해설:"):
        return strip_bullet(text.split(":", 1)[1])
    return ""


def simplify_question(text: str) -> str:
    text = normalize(re.sub(r"^\d+\.\s*", "", text))
    prefixes = [
        "사회보장기본법상 ",
        "사회보장수급권의 포기에 관하여 ",
        "건강보험법상 ",
        "건강보험법에 관하여 ",
        "국민건강보험법상 ",
        "국민연금법상 ",
        "산업재해보상보험법상 ",
        "고용보험법상 ",
    ]
    for prefix in prefixes:
        if text.startswith(prefix):
            text = text[len(prefix):]
    text = re.sub(r"^(사회보장기본법상|국민건강보험법상|국민연금법상|산업재해보상보험법상|고용보험법상)\s*", "", text)
    text = text.rstrip(".")
    return text


def extract_wrong_part(question_text: str, correction_text: str) -> str:
    question_text = simplify_question(question_text)
    correction_text = normalize(correction_text)

    for phrase in TRIGGER_PHRASES:
        if phrase in question_text:
            return f"'{phrase}'라고 한 부분이 틀렸다."

    token_pairs = [
        ("국무총리", "보건복지부장관"),
        ("고용노동부장관", "보건복지부장관"),
        ("대통령령", "보건복지부령"),
        ("보건복지부령", "대통령령"),
        ("국민건강보험공단", "건강보험심사평가원"),
        ("장해보상연금", "장해보상일시금"),
        ("3년", "2년"),
        ("법률", "정관"),
    ]
    for wrong, right in token_pairs:
        if wrong in question_text and right in correction_text:
            return f"'{wrong}'이라고 한 부분이 틀렸다."

    if "없다" in question_text and "있다" in correction_text:
        return "'없다'고 단정한 부분이 틀렸다."
    if "있다" in question_text and "없다" in correction_text:
        return "'있다'고 한 부분이 틀렸다."

    if len(question_text) <= 72:
        return f"'{question_text}'라고 한 부분이 틀렸다."
    return "지문의 주체·요건·효과를 잘못 적은 부분이 틀렸다."


def format_explanation(question_text: str, correction_text: str) -> str:
    correction_text = strip_bullet(correction_text).rstrip(".")
    wrong_part = extract_wrong_part(question_text, correction_text)
    return f"해설: 틀린 부분: {wrong_part} 정정: {correction_text}."


def main() -> int:
    helpers = runpy.run_path("scripts/replace_explanations_from_previous_docs.py")
    configs = helpers["CONFIGS"]
    social_cfg = next(cfg for cfg in configs if cfg["name"] == "social")
    load_candidates = helpers["load_candidates"]
    build_index = helpers["build_index"]
    choose_reference = helpers["choose_reference"]

    if not DOCX_PATH.exists():
        raise SystemExit(f"문서를 찾지 못했습니다: {DOCX_PATH}")
    if not PREVIOUS_WITH_BASIS.exists():
        raise SystemExit(f"근거 복원용 백업을 찾지 못했습니다: {PREVIOUS_WITH_BASIS}")

    backup_doc = backup_file(DOCX_PATH, "backup_before_fix_x_explanations")
    backup_pdf = backup_file(PDF_PATH, "backup_before_fix_x_explanations") if PDF_PATH.exists() else None

    basis_map = build_basis_map(PREVIOUS_WITH_BASIS)
    direct_x_map = build_direct_x_map(DIRECT_SOURCE_DOC) if DIRECT_SOURCE_DOC.exists() else {}
    candidates = load_candidates(social_cfg["reference"], social_cfg["stopwords"])
    keyword_index = build_index(candidates)

    doc = Document(DOCX_PATH)
    common_map = build_common_map(doc)

    current_q: int | None = None
    current_question_text = ""
    current_answer = ""
    updated = 0
    fallback_used = 0

    for p in doc.paragraphs:
        text = normalize(p.text)
        if not text:
            continue
        if re.match(r"^\d+\.\s", text):
            current_q = int(text.split(".", 1)[0])
            current_question_text = text
            current_answer = ""
            continue
        if text.startswith("정답:"):
            current_answer = text.split(":", 1)[1].strip()
            continue
        if current_q is None or current_answer != "X":
            continue
        if not (text.startswith("해설:") or text.startswith("공통해설(")):
            continue

        basis = basis_map.get(current_q, "")
        correction_text = MANUAL_CORRECTIONS.get(current_q, "")
        if not correction_text:
            correction_text = direct_x_map.get(current_q, "")
        if not correction_text:
            ref_text = choose_reference(
                current_question_text.split(". ", 1)[1].strip(),
                basis,
                candidates,
                keyword_index,
                social_cfg["stopwords"],
            )
            correction_text = strip_bullet(ref_text)
        if not correction_text:
            correction_text = resolve_current_explanation(text, common_map)
        if not correction_text:
            correction_text = "지문의 문언이 법정 기준과 다르므로 해당 조문의 주체·요건·효과를 정확한 문언으로 고쳐야 한다"
            fallback_used += 1

        p.text = format_explanation(current_question_text, correction_text)
        updated += 1

    doc.save(DOCX_PATH)

    pdf_helpers = runpy.run_path("scripts/optimize_pdfs_for_kindle_scribe.py")
    optimize_docx_for_kindle = pdf_helpers["optimize_docx_for_kindle"]
    export_pdf_with_pages = pdf_helpers["export_pdf_with_pages"]
    inspect_pdf = pdf_helpers["inspect_pdf"]

    with TemporaryDirectory() as tmpdir:
        temp_docx = Path(tmpdir) / DOCX_PATH.name
        optimize_docx_for_kindle(DOCX_PATH, temp_docx)
        export_pdf_with_pages(temp_docx, PDF_PATH)

    pages, width_in, height_in = inspect_pdf(PDF_PATH)
    print(f"backup_doc={backup_doc}")
    print(f"backup_pdf={backup_pdf}")
    print(f"updated_x_explanations={updated}")
    print(f"fallback_used={fallback_used}")
    print(f"pdf_pages={pages}")
    print(f"pdf_size_in={width_in:.2f}x{height_in:.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
