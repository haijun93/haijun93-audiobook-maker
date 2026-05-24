#!/usr/bin/env python3
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document

BASE_DIR = Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/260416")
BACKUP_DIR = BASE_DIR / "backup"

DOCX_FILES = [
    BASE_DIR / "[doc] social_insurance_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] labor_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] civil_law_OX_integrated_2026_v1.docx",
    BASE_DIR / "[doc] management_OX_integrated_2026_v1.docx",
]


@dataclass
class QuestionBlock:
    number: int
    question_idx: int
    explanation_idx: int | None
    explanation: str


def backup_file(path: Path) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"{path.stem}_backup_before_group_explanations_{ts}{path.suffix}"
    backup.write_bytes(path.read_bytes())
    return backup


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def parse_blocks(doc: Document) -> list[QuestionBlock]:
    blocks: list[QuestionBlock] = []
    current: QuestionBlock | None = None
    for idx, p in enumerate(doc.paragraphs):
        text = normalize(p.text)
        if re.match(r"^\d+\.\s", text):
            number = int(text.split(".", 1)[0])
            current = QuestionBlock(number=number, question_idx=idx, explanation_idx=None, explanation="")
            blocks.append(current)
        elif current and text.startswith("해설:"):
            current.explanation_idx = idx
            current.explanation = text.split(":", 1)[1].strip()
    return blocks


def find_runs(blocks: list[QuestionBlock]) -> list[list[QuestionBlock]]:
    runs: list[list[QuestionBlock]] = []
    i = 0
    while i < len(blocks):
        first = blocks[i]
        if not first.explanation:
            i += 1
            continue
        run = [first]
        j = i + 1
        while j < len(blocks):
            prev = blocks[j - 1]
            cur = blocks[j]
            if cur.number != prev.number + 1:
                break
            if not cur.explanation or cur.explanation != first.explanation:
                break
            run.append(cur)
            j += 1
        if len(run) >= 2:
            runs.append(run)
        i = j
    return runs


def apply_grouping(path: Path) -> tuple[Path, int]:
    backup = backup_file(path)
    doc = Document(path)
    blocks = parse_blocks(doc)
    runs = find_runs(blocks)

    for run in runs:
        start = run[0].number
        end = run[-1].number
        common_text = run[0].explanation
        first_para = doc.paragraphs[run[0].explanation_idx]
        first_para.text = f"공통해설(문항 {start}~{end}): {common_text}"
        for block in run[1:]:
            para = doc.paragraphs[block.explanation_idx]
            para.text = f"해설: 공통해설(문항 {start}~{end}) 참조"

    doc.save(path)
    return backup, len(runs)


def main() -> None:
    for path in DOCX_FILES:
        backup, groups = apply_grouping(path)
        print(f"{path.name}: backup={backup} groups={groups}")


if __name__ == "__main__":
    main()
