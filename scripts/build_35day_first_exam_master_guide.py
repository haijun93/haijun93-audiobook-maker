#!/usr/bin/env python3

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT_DIR = Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/260416")
DOCX_PATH = OUT_DIR / "[doc] 35day_first_exam_master_guide_2026.docx"
PDF_PATH = OUT_DIR / "[pdf] 35day_first_exam_master_guide_2026.pdf"


def set_font(run, name="Malgun Gothic", size=10.5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def configure_document(doc: Document) -> None:
    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    styles["Normal"].font.size = Pt(10.5)
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    section.page_width = Inches(6.19)
    section.page_height = Inches(8.26)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_font(run, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_font(run, size=10, bold=False)


def add_bullets(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(line)
        set_font(run)


def add_numbered(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(line)
        set_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        p = cell.paragraphs[0]
        run = p.add_run(header)
        set_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            run = p.add_run(value)
            set_font(run)


def build_doc() -> Document:
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        "35일 공인노무사 1차 마스터 가이드",
        "2026 시험 대비 / 노동법·사회보험법·민법·경영학 압축 회독 전략 / 2026-04-16",
    )

    doc.add_paragraph("")
    doc.add_paragraph("1. 이 자료의 목적", style="Heading 1")
    add_bullets(
        doc,
        [
            "시험일까지 35일 남은 시점에서, 이미 만들어 둔 4과목 OX 자료를 실제 점수로 연결하기 위한 운영 지침서다.",
            "핵심은 모든 내용을 새로 배우는 것이 아니라, 빈출·고난도·함정 포인트를 압축 회독하고 기출 적중률을 최대화하는 데 있다.",
            "학습의 기준은 60점 방어가 아니라 평균과 과락 리스크를 동시에 관리하는 '실전 합격형 점수 구조'에 둔다.",
        ],
    )

    doc.add_paragraph("2. 현재 자료 검증 상태", style="Heading 1")
    add_table(
        doc,
        ["과목", "기출 커버 상태", "해석", "메모"],
        [
            ["노동법", "658 / 660", "거의 완전", "복수정답·전원정답 성격 2문항 제외"],
            ["민법", "326 / 330", "매우 높음", "복수정답·전원정답 성격 4문항 제외"],
            ["경영학", "327 / 330", "매우 높음", "조합형 OCR 불안정 3문항만 별도 예외"],
            ["사회보험법", "328문항 기반 458 OX 포인트", "실전용으로 충분", "원문 선택지 1:1이 아니라 핵심 포인트 재구성형"],
        ],
    )
    add_bullets(
        doc,
        [
            "실전 학습 기준으로는 노동법·민법·경영학 모두 최근 기출 복구 수준이 매우 높다.",
            "사회보험법은 원문 선택지 재현본이라기보다, 2014~2025 기출을 중복 제거 후 핵심 포인트형 OX로 재구성한 자료다.",
        ],
    )

    doc.add_paragraph("3. 35일 전체 운영 원칙", style="Heading 1")
    add_numbered(
        doc,
        [
            "매일 4과목을 모두 건드리되, 시간 배분은 노동법 > 민법 > 사회보험법 > 경영학 순으로 둔다.",
            "새로운 강의나 새로운 교재를 넓게 펼치지 말고, 현재 만든 OX 통합본과 이 가이드만 반복 회독의 중심에 둔다.",
            "틀린 문항은 '해설 완독'보다 '핵심-기억-#두문자-암기비법-함께'를 기준으로 재회독한다.",
            "박스형·케이스형·조합형 문제는 정답만 맞히지 말고, 틀린 선택지가 왜 틀렸는지까지 말로 설명할 수 있어야 한다.",
            "최신 개정·시행령·위원회·숫자·기간·절차는 마지막 14일 동안 매일 반복 노출한다.",
        ],
    )

    doc.add_paragraph("4. 과목별 점수 전략", style="Heading 1")
    add_table(
        doc,
        ["과목", "현실 목표", "공부 원칙", "막판 포인트"],
        [
            ["노동법", "70~80+", "조문·시행령·위원회·판례를 한 세트로 회독", "부속법령, 절차, 숫자, 최신 판례"],
            ["사회보험법", "75~85+", "용어정의·위원회·급여·보험료·심사절차 비교 학습", "숫자, 기간, 신청/심사/재심 흐름"],
            ["민법", "65~75+", "조문 기본기 + 중요 판례 + 박스형 조합형 적응", "보증·대리·채권총론·계약·케이스형"],
            ["경영학", "60~70+", "인사/조직 + 전략 고득점, 재무/회계는 빈출만 압축", "재무 계산 핵심, 조직/전략/마케팅 빈출"],
        ],
    )

    doc.add_paragraph("5. 35일 세부 일정", style="Heading 1")
    add_table(
        doc,
        ["구간", "핵심 목표", "과목별 중점"],
        [
            ["D-35 ~ D-29", "전체 구조 재정렬", "노동법·사회보험법 비교표 완독 / 민법 총칙·채권총론 / 경영학 인사·조직"],
            ["D-28 ~ D-22", "빈출 파트 1회 완주", "노동법1 부속법령 / 사회보험 급여·보험료 / 민법 채권각론 / 경영학 전략·마케팅"],
            ["D-21 ~ D-15", "기출형 사고 적응", "2019~2025 기출 회독 / 박스형·케이스형·조합형 집중"],
            ["D-14 ~ D-10", "약점 수리", "노동법 위원회·절차 / 사회보험 숫자·위원회 / 민법 판례 문구 / 경영학 재무·회계 빈출"],
            ["D-9 ~ D-5", "실전 점수화", "과목별 40문항·25문항 시간 제한 풀이 / 오답만 재회독"],
            ["D-4 ~ D-2", "최종 암기", "비교표·#두문자·숫자·기간·절차만 회독"],
            ["D-1", "컨디션 조절", "새 내용 금지 / 최종 오답노트와 비교표 1회독"],
        ],
    )

    doc.add_paragraph("6. 하루 운영 템플릿", style="Heading 1")
    add_table(
        doc,
        ["시간대", "내용", "운영 방식"],
        [
            ["1교시", "노동법", "조문·판례·위원회 표 회독 + OX 60~80문항"],
            ["2교시", "민법", "케이스·박스형 중심 OX 50~70문항 + 판례 확인"],
            ["3교시", "사회보험법", "숫자·절차·위원회 비교표 후 OX 50문항"],
            ["4교시", "경영학", "인사/조직·전략 고정 + 재무/회계 빈출 30~40문항"],
            ["마감", "오답 정리", "틀린 문항만 다시 보고 #두문자와 암기비법으로 압축"],
        ],
    )
    add_bullets(
        doc,
        [
            "노동법은 매일, 민법은 거의 매일, 사회보험법·경영학은 번갈아 더 길게 배치하되 완전히 비우는 날은 만들지 않는다.",
            "하루 총량이 무너지면 '전 범위 얕게'보다 '노동법 + 민법 + 사회보험법 숫자표'만이라도 반드시 수행한다.",
        ],
    )

    doc.add_paragraph("7. 과목별 최우선 회독 순서", style="Heading 1")
    add_table(
        doc,
        ["과목", "1순위", "2순위", "3순위", "막판 체크"],
        [
            ["노동법", "근로기준법 / 노동조합법", "부속법령 / 노동위원회법 / 근참법", "최신 판례 / 시행령 / 시행규칙", "위원회·기한·구성"],
            ["사회보험법", "고용보험·산재보험·징수법", "건강보험·국민연금", "사회보장기본법", "위원회·심사·재심·기간"],
            ["민법", "총칙·대리·의사표시", "채권총론 / 상계 / 변제 / 소멸시효", "매매·임대차·도급·불법행위", "박스형·케이스형 조합"],
            ["경영학", "인사/조직", "경영전략·마케팅", "운영관리", "재무/회계 빈출 계산"],
        ],
    )

    doc.add_paragraph("8. 함정 회피 체크리스트", style="Heading 1")
    add_bullets(
        doc,
        [
            "노동법: 법률 조문만 보지 말고 시행령·시행규칙까지 함께 본다. 위원회 문제는 소속·구성·의결 정족수를 묶어 외운다.",
            "사회보험법: 급여 종류, 지급요건, 기간, 보험료율, 심사/재심사 기한을 비교표로 보지 않으면 헷갈리기 쉽다.",
            "민법: 옳은 판례 하나보다 틀린 선택지 포인트를 먼저 본다. 특히 조합형 문항은 ㄱ·ㄴ·ㄷ 각각 독립 판단이 핵심이다.",
            "경영학: 인사/조직은 개념 비교, 전략은 프레임워크 비교, 재무/회계는 식 자체보다 어떤 공식을 고르는지부터 점검한다.",
        ],
    )

    doc.add_paragraph("9. 마지막 7일 체크리스트", style="Heading 1")
    add_numbered(
        doc,
        [
            "노동법 위원회·근로자 대표·노조 관련 숫자와 절차를 하루 1회 확인한다.",
            "사회보험법 위원회·심사·재심·보험급여·보험료 숫자표를 매일 본다.",
            "민법은 박스형·케이스형 문제만 따로 묶어 다시 본다.",
            "경영학은 인사/조직·전략을 고정 득점 파트로 유지하고 재무/회계는 빈출 계산형만 확인한다.",
            "새 교재나 새 문제보다 지금까지 틀린 문제와 비교표를 반복한다.",
        ],
    )

    doc.add_paragraph("10. 예외문항 메모", style="Heading 1")
    add_bullets(
        doc,
        [
            "노동법: 2021년 26번, 2024년 43번은 복수정답/전원정답 성격 예외문항으로 별도 취급했다.",
            "민법: 2021년 70번, 2022년 51번·72번, 2025년 26번은 복수정답/전원정답 성격 예외문항으로 별도 취급했다.",
            "경영학: 현재 OCR 복구 기준으로 2020년 123번, 2021년 110번·123번은 원문 조합형 복원이 덜 안정적이므로 최종 원문 재확인 권장 구간으로 표시한다.",
        ],
    )

    doc.add_paragraph("11. 바로 실행할 오늘의 루틴", style="Heading 1")
    add_numbered(
        doc,
        [
            "노동법 비교표 20분 + OX 60문항",
            "민법 박스형·케이스형 40문항",
            "사회보험법 숫자·위원회 표 20분 + OX 40문항",
            "경영학 인사/조직 30문항 + 재무/회계 빈출 계산 10문항",
            "틀린 문항만 다시 보고 #두문자와 암기비법으로 15분 압축 정리",
        ],
    )

    return doc


def export_pdf(docx_path: Path, pdf_path: Path) -> None:
    applescript = f"""
set inputPath to POSIX file "{docx_path.as_posix()}"
set pdfPath to POSIX file "{pdf_path.as_posix()}"

tell application "Pages"
\tactivate
\tset docRef to open inputPath
\tdelay 2
\texport docRef to pdfPath as PDF
\tclose docRef saving no
end tell
"""
    import subprocess

    subprocess.run(["osascript", "-e", applescript], check=True)


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(DOCX_PATH)
    export_pdf(DOCX_PATH, PDF_PATH)
    print(DOCX_PATH)
    print(PDF_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
