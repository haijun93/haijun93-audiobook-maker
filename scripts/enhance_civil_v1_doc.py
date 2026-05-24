#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from kiwipiepy import Kiwi


DEFAULT_INPUT = Path(
    "/Users/hyeokjunkong/Desktop/1차 시험/#STD/260416/civil_law_OX_integrated_2026_v1.docx"
)
KIWI = Kiwi()

SECTION_ORDER = [
    "총칙·권리능력·법원",
    "의사표시·법률행위·대리",
    "물건·물권",
    "채권·계약·불법행위",
    "법인",
    "친족·상속",
    "기타",
]

SECTION_NOTES = {
    "총칙·권리능력·법원": "민법 총칙은 용어, 법원, 권리능력, 제한능력자, 소멸시효처럼 기본 개념을 한 단어씩 바꿔 내는 문제가 많다. 조문 순서와 판례 결론을 함께 붙여서 기억하는 편이 좋다.",
    "의사표시·법률행위·대리": "의사표시와 대리는 요건과 효과를 미세하게 바꾸는 함정이 많다. 무효·취소·추인·철회·표현대리·대리권 남용을 한 묶음으로 비교해야 한다.",
    "물건·물권": "물권 파트는 공시, 등기, 점유, 선의취득, 용익물권의 요건을 조금씩 비틀어 출제하는 경우가 많다. 한 요소라도 빠지면 결론이 달라진다.",
    "채권·계약·불법행위": "최근 기출의 중심 파트다. 채권총론 10문제, 채권각론 15문제로 무게가 실려 있고, 사례형과 박스형도 가장 자주 나온다. 주체, 항변, 해제, 손해배상, 담보책임, 불법행위 성립요건을 함께 정리해야 한다.",
    "법인": "법인은 출제량이 많지는 않아도 청산, 대표권 제한, 불법행위책임처럼 틀리기 쉬운 지점이 반복된다. 조문과 판례를 세트로 정리하는 편이 안전하다.",
    "친족·상속": "친족·상속은 출제 비중은 상대적으로 낮지만, 개념을 정확히 알고 있으면 쉽게 맞힐 수 있는 파트다. 용어와 순위, 효과를 정확히 기억하는 것이 중요하다.",
    "기타": "민법 기타 파트는 여러 장의 논점이 섞여 있어 낯설게 느껴질 수 있다. 같은 쟁점에서 누가 무엇을 주장할 수 있는지 중심으로 보는 것이 효율적이다.",
}

SECTION_GUIDES = {
    "총칙·권리능력·법원": {
        "key": "민법 총칙은 법원, 권리능력, 제한능력자, 소멸시효처럼 기본 개념의 정의와 요건을 정확히 구분하는 것이 핵심이다.",
        "memory": "총칙 문제는 '개념-요건-효과' 세 칸으로 나눠 정리하면 조문과 판례가 함께 기억된다.",
        "related": "관습법, 신의칙, 소멸시효는 다른 파트에서도 반복 연결되므로 총칙에서 기준을 먼저 고정해 두는 편이 좋다.",
    },
    "의사표시·법률행위·대리": {
        "key": "의사표시, 무효·취소, 추인, 철회, 표현대리, 대리권 남용은 요건과 효과를 세트로 구분하는 것이 핵심이다.",
        "memory": "이 파트는 '누가 말했고 어떤 하자가 있으며 그 결과가 무엇인지' 순서로 읽으면 정리가 쉽다.",
        "related": "표현대리, 무권대리, 착오·사기·강박은 서로 요건을 바꿔 묻는 경우가 많으므로 비교표로 묶어 둔다.",
    },
    "물건·물권": {
        "key": "공시, 등기, 점유, 선의취득, 용익물권의 성립 요건과 효과를 정확히 구분하는 것이 핵심이다.",
        "memory": "물권은 '무엇이 공시 방법인지'와 '누가 어떤 권리를 취득하는지'를 먼저 체크하면 결론이 빨리 잡힌다.",
        "related": "점유와 등기, 선의취득과 무권리자 처분, 지상권과 건물매수청구권은 함께 비교해 두는 편이 좋다.",
    },
    "채권·계약·불법행위": {
        "key": "채권총칙과 각론은 항변, 해제, 손해배상, 담보책임, 불법행위 성립요건처럼 효과가 달라지는 포인트를 정확히 구분해야 한다.",
        "memory": "계약 문제는 '누가 먼저 이행해야 하는지'와 '계약이 유지되는지 종료되는지'를 먼저 표시하면 풀기 쉬워진다.",
        "related": "동시이행항변권, 해제권, 손해배상 범위, 보증채무, 불법행위 책임은 사례형에서도 반복되므로 함께 연결해 본다.",
    },
    "법인": {
        "key": "법인은 대표권, 청산, 기관 행위, 불법행위책임의 구조를 정확히 구분하는 것이 핵심이다.",
        "memory": "법인 파트는 '누가 법인을 대표하는지'와 '청산 단계인지'를 먼저 체크하면 오답이 줄어든다.",
        "related": "대표권 제한, 이사회 소집, 청산사무, 불법행위책임은 기관 구조와 함께 정리해 두는 편이 좋다.",
    },
    "친족·상속": {
        "key": "친족·상속은 순위, 범위, 효과를 정확히 기억하는 것이 가장 중요하다.",
        "memory": "순위가 나오는 파트는 표로 외우고, 효과는 '누가 무엇을 취득하는지'로 짧게 정리하면 기억이 쉽다.",
        "related": "상속순위, 유류분, 상속승인·포기처럼 절차와 효과가 달라지는 부분을 함께 봐 두면 안정적이다.",
    },
    "기타": {
        "key": "낯선 쟁점일수록 주체, 요건, 효과를 먼저 잡는 것이 핵심이다.",
        "memory": "문제가 길어도 당사자와 시간 순서를 먼저 표시하면 논점이 빨리 정리된다.",
        "related": "같은 쟁점에서 누가 어떤 권리를 행사할 수 있는지 중심으로 연결해서 보면 회독이 편해진다.",
    },
}

SECTION_EXAMPLES = {
    "총칙·권리능력·법원": "예를 들어 법률에 규정이 없을 때 곧바로 판례를 찾는 것이 아니라, 먼저 관습법이 있는지 보고 그다음 조리로 가는 순서를 떠올리면 총칙 문제가 쉬워진다.",
    "의사표시·법률행위·대리": "예를 들어 대리권 없이 계약을 맺은 사람이 있어도 본인이 나중에 추인하면 유효가 될 수 있고, 표현대리는 겉으로 보이는 권한이 핵심이 된다.",
    "물건·물권": "예를 들어 부동산은 등기, 동산은 인도처럼 공시방법이 다르다는 점을 먼저 잡으면 물권 문제를 훨씬 빨리 풀 수 있다.",
    "채권·계약·불법행위": "예를 들어 매매계약에서 한쪽이 먼저 이행해야 하는지, 서로 동시에 이행해야 하는지가 달라지면 동시이행항변권 결론도 달라진다.",
    "법인": "예를 들어 이사가 회사를 대표하는지, 청산인이 대표하는지에 따라 법률관계가 달라지므로 현재 단계가 평상시인지 청산 중인지 먼저 봐야 한다.",
    "친족·상속": "예를 들어 상속은 누가 먼저 상속인이 되는지 순위부터 잡고, 그다음 유류분이나 승인·포기 효과를 보는 식으로 접근하면 정리가 쉽다.",
    "기타": "예를 들어 당사자, 시간 순서, 효과를 먼저 적어 보면 복잡한 사례형 지문도 훨씬 단순해진다.",
}

ISSUE_SPECIFIC_GUIDES = (
    (
        re.compile(r"(미분리 과실|공시 방법)"),
        {
            "key": "관습법상 미분리 과실 문제는 관행의 존재만이 아니라 그 관행을 외부에 어떻게 공시하는지가 함께 인정되는지를 보는 것이 핵심이다.",
            "memory": "미분리 과실은 '과실도 공시가 필요하다'는 한 줄로 먼저 잡아 두면 정리가 빠르다.",
            "related": "관습법상 공시 방법 인정 여부는 명문 규정이 없는 영역에서 법원이 어떻게 법원성을 보충하는지와 연결해 이해하면 좋다.",
        },
    ),
    (
        re.compile(r"(종중|공동 선조와 성과 본을 같이 하는 후손인 여성)"),
        {
            "key": "종중 파트는 종중 구성원의 범위를 전통적 관념이 아니라 현재 판례 기준에 따라 파악하는 것이 핵심이다.",
            "memory": "종중은 '공동 선조 + 성년 + 성과 본'이라는 기준틀로 먼저 외우면 정리가 쉽다.",
            "related": "관습법도 전체 법질서에 부합해야 효력이 인정되므로, 종중 판례는 관습법 파트와 함께 묶어 두는 편이 좋다.",
        },
    ),
    (
        re.compile(r"(관습법|법원성)"),
        {
            "key": "관습법 문제는 성립 요건과 법원성 인정 범위를 정확히 구분하는 것이 핵심이다.",
            "memory": "관습법은 '반복된 관행 + 법적 확신 + 전체 법질서 부합' 세 요소로 외우면 정리가 쉽다.",
            "related": "조리와 관습법의 적용 순서, 종중 관련 판례는 법원 파트와 함께 묶어 두는 편이 좋다.",
        },
    ),
    (
        re.compile(r"(조리|법률에 규정이 없으면)"),
        {
            "key": "법원의 적용 순서는 법률, 관습법, 조리의 단계라는 점을 정확히 구분하는 것이 핵심이다.",
            "memory": "민사 법원 순서는 '법률 먼저, 없으면 관습법, 끝까지 없으면 조리'로 한 줄 정리하면 된다.",
            "related": "관습법은 독자적 법원이고, 조리는 최후 보충 규범이라는 점을 함께 비교해 두면 좋다.",
        },
    ),
    (
        re.compile(r"(미성년자|제한능력자|법정대리인|성년이 된 후|증여|경매 절차|철회|추인|상대방의 물음|취소할 수 없다|취소할 수 있다)"),
        {
            "key": "제한능력자 파트는 누가 동의해야 하는지와 취소·추인·철회의 효과를 정확히 구분하는 것이 핵심이다.",
            "memory": "이 파트는 '동의-취소-추인' 세 단계로 나눠 외우면 헷갈림이 줄어든다.",
            "related": "상대방의 철회권, 단독으로 할 수 있는 행위, 법정대리인의 동의 필요 여부를 함께 비교해 둔다.",
        },
    ),
    (
        re.compile(r"(실종선고|부재자|재산관리인)"),
        {
            "key": "실종선고 파트는 선고 전 관리와 선고 후 효과, 취소 후 법률관계를 구분하는 것이 핵심이다.",
            "memory": "실종선고는 '관리-선고-취소 후 반환' 순서로 정리하면 구조가 잘 보인다.",
            "related": "선의의 제3자 보호와 현존이익 반환 범위는 실종선고 취소 효과와 함께 묶어 둔다.",
        },
    ),
    (
        re.compile(r"(신의성실|신의칙)"),
        {
            "key": "신의칙은 적용 요건과 제한 효과를 구체적 사실관계와 연결해 이해하는 것이 핵심이다.",
            "memory": "신의칙은 '형평·신뢰 보호'라는 큰 틀과 대표 판례를 같이 붙여 두면 오래 남는다.",
            "related": "권리남용, 사정변경, 금반언과 함께 비교하면 적용 범위가 더 선명해진다.",
        },
    ),
    (
        re.compile(r"(의사표시|착오|사기|강박|비진의|통정허위표시)"),
        {
            "key": "의사표시는 하자의 종류별 요건과 효과를 구분하는 것이 핵심이다.",
            "memory": "의사표시는 '무효인지 취소인지'부터 먼저 분류하면 정리가 빨라진다.",
            "related": "착오, 사기·강박, 통정허위표시는 제3자 보호 여부까지 함께 비교해 둔다.",
        },
    ),
    (
        re.compile(r"(표현대리|무권대리|대리권)"),
        {
            "key": "대리 파트는 기본대리권 존재, 권한 범위, 상대방 보호 요건을 정확히 구분하는 것이 핵심이다.",
            "memory": "표현대리는 '기본대리권-외관-상대방 신뢰' 순서로 외우면 구조가 잘 보인다.",
            "related": "무권대리, 표현대리, 대리권 남용은 모두 본인 귀속 여부를 중심으로 비교하는 것이 좋다.",
        },
    ),
    (
        re.compile(r"(소멸시효|제척기간)"),
        {
            "key": "시효 파트는 기산점, 중단·정지, 완성 효과를 구분하는 것이 핵심이다.",
            "memory": "시효는 '언제부터 세는지'를 먼저 잡으면 절반은 풀린다.",
            "related": "중단 사유, 원용 주체, 제척기간과의 차이를 함께 비교해 둔다.",
        },
    ),
    (
        re.compile(r"(보증채무|보증인)"),
        {
            "key": "보증채무는 주채무와의 관계, 부종성, 시효, 보증인 책임 범위를 함께 보는 것이 핵심이다.",
            "memory": "보증은 '주채무를 따라가는가'를 먼저 체크하면 정리가 쉽다.",
            "related": "연대보증, 사정변경, 보증채무의 범위와 연체이율 문제를 같이 묶어 둔다.",
        },
    ),
    (
        re.compile(r"(건물매수청구권|지상물매수청구권|임대차|건물철거)"),
        {
            "key": "건물매수청구권은 행사 요건, 행사 효과, 대금·인도 관계를 정확히 구분하는 것이 핵심이다.",
            "memory": "이 파트는 '행사 가능 여부-매매 성립 효과-대금·인도 문제' 순서로 보면 정리가 쉽다.",
            "related": "무허가 건물, 경제적 가치, 지연손해금, 부당이득 반환 문제를 함께 비교해 둔다.",
        },
    ),
    (
        re.compile(r"(동시이행항변권|상계)"),
        {
            "key": "동시이행항변권과 상계는 항변 가능 여부와 자동채권·수동채권의 관계를 구분하는 것이 핵심이다.",
            "memory": "상계는 '자동채권이 상계 가능한 상태인지'부터 먼저 보면 풀기 쉽다.",
            "related": "이행지체, 변제충당, 공탁과 함께 연결해 두면 채권총론 파트가 정리된다.",
        },
    ),
    (
        re.compile(r"(해제|위험의 이전|이행불능|매매계약)"),
        {
            "key": "계약 해제 파트는 해제권 발생 요건, 위험부담, 해제의 효과를 구분하는 것이 핵심이다.",
            "memory": "계약 문제는 '계약이 유지되는지 끝나는지'를 먼저 표시하면 구조가 빨리 잡힌다.",
            "related": "해제약관, 이행지체, 위험이전, 담보책임은 함께 비교해 두는 편이 좋다.",
        },
    ),
    (
        re.compile(r"(불법행위|부당이득|과실상계|입증책임)"),
        {
            "key": "불법행위와 부당이득은 성립 요건, 입증책임, 효과를 정확히 구분하는 것이 핵심이다.",
            "memory": "불법행위는 '위법성-고의·과실-인과관계', 부당이득은 '이익-손실-법률상 원인 없음'으로 외우면 깔끔하다.",
            "related": "손해배상 범위, 과실상계, 사용자책임, 증명책임 분배까지 같이 묶어 두면 좋다.",
        },
    ),
)

SECTION_MNEMONICS = {
    "총칙·권리능력·법원": {"tag": "#개요효", "tip": "개념-요건-효과 세 칸으로 나눠 적으면 민법 총칙이 가장 안정적으로 정리된다."},
    "의사표시·법률행위·대리": {"tag": "#하주효", "tip": "하자-주체-효과를 먼저 구분하면 긴 문장도 빠르게 정리된다."},
    "물건·물권": {"tag": "#공취권", "tip": "공시-취득-권리 구조로 묶으면 물권 파트가 훨씬 선명해진다."},
    "채권·계약·불법행위": {"tag": "#항해손", "tip": "항변-해제-손해배상 순서로 외우면 채권법 뼈대가 잡힌다."},
    "법인": {"tag": "#대청기", "tip": "대표-청산-기관 구조를 먼저 외우고 세부 판례를 덧붙인다."},
    "친족·상속": {"tag": "#순범효", "tip": "순위-범위-효과를 표로 정리하면 상속 파트는 오래 남는다."},
    "기타": {"tag": "#주요효", "tip": "주체-요건-효과를 먼저 적는 습관이 낯선 사례형에도 가장 잘 통한다."},
}

ISSUE_MNEMONICS = (
    (re.compile(r"(미분리 과실|공시 방법)"), {"tag": "#관공", "tip": "관습법상 공시 인정 문제라고 묶어 외우면 미분리 과실 포인트가 쉽게 남는다."}),
    (re.compile(r"(종중|공동 선조와 성과 본을 같이 하는 후손인 여성)"), {"tag": "#공성본", "tip": "공동선조-성년-성과 본의 기준으로 종중 구성원 범위를 정리해 둔다."}),
    (re.compile(r"(관습법|법원성)"), {"tag": "#반확질", "tip": "반복된 관행-법적 확신-전체 법질서 부합의 세 요소를 순서대로 외운다."}),
    (re.compile(r"(조리|법률에 규정이 없으면)"), {"tag": "#법관조", "tip": "법률-관습법-조리의 적용 순서를 한 줄로 반복해서 읽으면 잘 안 헷갈린다."}),
    (re.compile(r"(미성년자|제한능력자|법정대리인|성년이 된 후|증여|경매 절차|철회|추인|상대방의 물음|취소할 수 없다|취소할 수 있다)"), {"tag": "#동취추", "tip": "동의-취소-추인의 순서로 외우고, 철회권은 상대방 쪽에 붙여 따로 표시한다."}),
    (re.compile(r"(실종선고|부재자|재산관리인)"), {"tag": "#관리선반", "tip": "관리-실종선고-반환 순서로 구조를 잡으면 실종선고 효과가 오래 남는다."}),
    (re.compile(r"(신의성실|신의칙)"), {"tag": "#형신", "tip": "형평-신뢰 보호라고 짧게 외우고 대표 판례 하나를 바로 연결한다."}),
    (re.compile(r"(의사표시|착오|사기|강박|비진의|통정허위표시)"), {"tag": "#무취제", "tip": "무효인지 취소인지, 제3자 보호가 있는지 두 칸으로 나눠 외운다."}),
    (re.compile(r"(표현대리|무권대리|대리권)"), {"tag": "#기외신", "tip": "기본대리권-외관-신뢰의 순서로 표현대리 구조를 반복해서 읽는다."}),
    (re.compile(r"(소멸시효|제척기간)"), {"tag": "#기중완", "tip": "기산점-중단/정지-완성효과 순서로 보면서 시효 문제를 정리한다."}),
    (re.compile(r"(보증채무|보증인)"), {"tag": "#주부범", "tip": "주채무-부종성-범위 순서로 보면 보증채무의 핵심이 빠르게 잡힌다."}),
    (re.compile(r"(건물매수청구권|지상물매수청구권|임대차|건물철거)"), {"tag": "#행매대", "tip": "행사 요건-매매 성립-대금·인도 문제를 순서대로 외운다."}),
    (re.compile(r"(동시이행항변권|상계)"), {"tag": "#항자상", "tip": "항변 가능 여부-자동채권 상태-상계 가능성을 차례대로 체크한다."}),
    (re.compile(r"(해제|위험의 이전|이행불능|매매계약)"), {"tag": "#해위효", "tip": "해제 발생-위험 이전-효과를 한 묶음으로 외우면 계약 파트가 빨라진다."}),
    (re.compile(r"(불법행위|부당이득|과실상계|입증책임)"), {"tag": "#위고인", "tip": "위법성-고의과실-인과관계를 먼저 외우고, 부당이득은 별도로 비교해서 본다."}),
)

GENERIC_EXPLANATION_PATTERNS = (
    r"같은 기출문항에서 옳은 진술로 기능하는 문장을 그대로 OX화한 것이다\.\s*",
    r"민법 총칙은.+?기억하는 편이 좋다\.\s*",
    r"민법은 조문 문언을 한두 단어 바꾸는 기본형 함정도 꾸준히 출제된다\.\s*",
    r"민법은 요건 하나가 빠지면 효과가 완전히 달라지므로 요건과 법률효과를 세트로 정리해야 한다\.\s*",
    r"사례형 문제는 사실관계 한 줄 차이로 결론이 달라지므로 당사자 지위와 시간 순서를 먼저 잡는 것이 좋다\.\s*",
)

FOCUS_STOPWORDS = {
    "민법", "판례", "조문", "문제", "지문", "설명", "기준", "경우", "법률", "내용",
    "당사자", "효과", "요건",
}

MANUAL_X_RULES = (
    (
        re.compile(r"민사에 관하여 법률에 규정이 없으면 조리에 의하고 조리가 없으면 관습법에 의한다"),
        "정답은 X입니다. 민사에 관하여 법률에 규정이 없으면 먼저 관습법에 의하고, 관습법이 없으면 조리에 의한다.",
    ),
    (
        re.compile(r"성년자인지에 대한 거래 상대방의 물음에 단순히 대답하지 않은 경우"),
        "정답은 X입니다. 단순히 성년 여부를 묻는 질문에 답하지 않은 것만으로는 사술에 의한 행위가 되지 않으므로, 곧바로 취소권이 배제되는 것은 아니다.",
    ),
    (
        re.compile(r"법원의 허가를 받아 한 건물의 처분 행위는 甲이 실종선고를 받게 되면 그 효력을 잃는다"),
        "정답은 X입니다. 부재자의 재산관리인이 법원의 허가를 받아 한 처분행위는 뒤에 실종선고가 선고되더라도 당연히 효력을 잃지 않는다.",
    ),
)

DIRECT_REPLACEMENTS = {
    "법률 행 위": "법률행위",
    "법률 행위를": "법률행위를",
    "법률 행위의": "법률행위의",
    "법률 행위": "법률행위",
    "채무 자": "채무자",
    "채권 자": "채권자",
    "법정 대리인": "법정대리인",
    "재판 외": "재판외",
    "양 도": "양도",
    "한정 후견": "한정후견",
    "성년 후견": "성년후견",
    "가정 법원": "가정법원",
    "소멸 시효": "소멸시효",
    "제척 기간": "제척기간",
    "부동산 매매 계약": "부동산 매매계약",
    "상당 한": "상당한",
    "하여야한다": "하여야 한다",
    "받아야한다": "받아야 한다",
    "될 수 없다": "될 수 없다",
    "법원(法院)": "법원",
    "민법상 ": "민법상 ",
}

REGEX_REPLACEMENTS = (
    (r"표현\s*대리", "표현대리"),
    (r"부진정\s*연대\s*책임", "부진정연대책임"),
    (r"채권자\s*취소권", "채권자취소권"),
    (r"채권자\s*대위권", "채권자대위권"),
    (r"선의\s*취득", "선의취득"),
    (r"동시\s*이행\s*항변권", "동시이행항변권"),
    (r"이행\s*불능", "이행불능"),
    (r"이행\s*지체", "이행지체"),
    (r"손해\s*배상", "손해배상"),
    (r"하자\s*담보\s*책임", "하자담보책임"),
    (r"예약\s*완결권", "예약완결권"),
    (r"건물\s*매수\s*청구권", "건물매수청구권"),
    (r"과실\s*상계", "과실상계"),
    (r"재산\s*관리인", "재산관리인"),
    (r"실종\s*선고", "실종선고"),
    (r"소유권\s*이전\s*등기", "소유권이전등기"),
    (r"지연\s*손해금", "지연손해금"),
    (r"형성권", "형성권"),
    (r"소급\s*효", "소급효"),
    (r"철회", "철회"),
    (r"해제", "해제"),
    (r"제\s*(\d+)\s*조", r"제\1조"),
)

LAW_URLS = [
    "https://law.go.kr/LSW/LsiJoLinkP.do?docType=JO&joNo=000800000&languageType=KO&lsNm=%EB%AF%BC%EB%B2%95&paras=1",
]

TRAP_MEMORY_ROWS = [
    (
        "법률·관습법·조리 순서",
        "민사에 법률 규정이 없으면 관습법, 그다음 조리 순서로 간다.",
        "#관조리",
        "관습법이 조리보다 앞선다는 점만 확실히 고정하면 기본 순서형 문제는 거의 틀리지 않는다.",
    ),
    (
        "제한능력자 효과",
        "동의가 필요한 행위, 취소권자, 추인, 확답촉구를 한 세트로 봐야 한다.",
        "#동취추확",
        "동의-취소-추인-확답촉구 순서로 읽고, 누가 행사하는지만 옆에 붙여 외우면 효율적이다.",
    ),
    (
        "표현대리",
        "기본대리권 존재, 권한 넘은 표현대리, 상대방의 정당한 이유를 구분해야 한다.",
        "#기외정",
        "기본권-외관-정당한 이유 세 단어를 연결해 외우면 표현대리 요건이 깔끔하게 정리된다.",
    ),
    (
        "시효 중단·정지·완성",
        "기산점, 중단사유, 완성 후 효과를 따로 보지 않으면 상계·승인 문제에서 흔들리기 쉽다.",
        "#기중완",
        "기산-중단-완성 3단 구조를 먼저 외운 뒤 사례에 끼워 넣는 연습을 하면 케이스형 대응력이 좋아진다.",
    ),
    (
        "보증채무·연대채무",
        "보증채무와 연대채무는 책임 구조와 부종성, 시효, 최고·검색 항변 등에서 차이가 난다.",
        "#보연부시",
        "보증-연대-부종성-시효를 표로 나눠 놓으면 비슷한 문장을 섞어도 구별이 쉬워진다.",
    ),
    (
        "해제·위험부담·이행불능",
        "채무불이행 유형과 계약 해제 효과, 위험부담 귀속을 함께 물어보는 지문이 반복된다.",
        "#해위이",
        "해제-위험부담-이행불능 순서로 연결하고, 누가 손해를 부담하는지 중심으로 기억하면 좋다.",
    ),
    (
        "임대차·건물매수청구권",
        "건물 소유 목적 토지임대차, 건물매수청구권 행사 요건, 대금과 인도 거절 관계가 자주 출제된다.",
        "#토건매",
        "토지-건물-매수청구권 흐름으로 그림을 그리며 외우면 판례형 지문도 훨씬 덜 막힌다.",
    ),
]

COMPARE_TABLE_ROWS = [
    ("법원 적용 순서", "법률 → 관습법 → 조리", "민사에 법률 규정이 없을 때 적용 순서를 묻는 기본 함정"),
    ("무효·취소·해제", "원시적 효력 부정 / 취소 가능 / 유효한 계약의 장래·소급 해소", "효과가 비슷해 보여도 출발점이 다르다"),
    ("제한능력자·무권대리·표현대리", "취소권 / 본인 책임 부정 / 외관 책임 인정", "주체와 효과를 함께 비교해야 한다"),
    ("이행지체·이행불능·불완전이행", "지연 / 이행 자체 불가 / 흠 있는 이행", "손해배상과 해제 연결 구조가 다르다"),
    ("보증채무·연대채무", "부종성·항변권 / 독립 책임 구조", "최고·검색의 항변과 시효를 같이 비교하면 좋다"),
    ("매도인 담보책임·채무불이행 책임", "하자·권리 흠결 / 일반 채무불이행", "요건과 효과를 한 줄로 구분해야 한다"),
]


@dataclass
class CivilQuestion:
    number: int
    section: str
    statement: str
    answer: str
    basis: str
    explanation: str
    source: str


def backup_file(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}_backup_before_enhance_{timestamp}{path.suffix}")
    backup.write_bytes(path.read_bytes())
    return backup


def normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = text.replace("•", "·").replace("ㆍ", "·")
    text = text.replace("“", '"').replace("”", '"').replace("’", "'").replace("‘", "'")
    return re.sub(r"\s+", " ", text).strip()


def needs_spacing(text: str) -> bool:
    compact = re.sub(r"[^가-힣]", "", text)
    if " " in text:
        return False
    return len(compact) >= 12


def space_text(text: str) -> str:
    updated = normalize_text(text)
    if not updated:
        return updated
    if needs_spacing(updated):
        updated = KIWI.space(updated)
    for before, after in DIRECT_REPLACEMENTS.items():
        updated = updated.replace(before, after)
    for pattern, replacement in REGEX_REPLACEMENTS:
        updated = re.sub(pattern, replacement, updated)
    updated = re.sub(r"\s+([.,;:])", r"\1", updated)
    updated = re.sub(r"\(\s+", "(", updated)
    updated = re.sub(r"\s+\)", ")", updated)
    updated = re.sub(r"\s{2,}", " ", updated)
    return updated.strip()


def parse_existing_doc(path: Path) -> list[CivilQuestion]:
    doc = Document(str(path))
    questions: list[CivilQuestion] = []
    section = ""
    current: dict[str, str | int] | None = None
    in_body = False

    def flush() -> None:
        nonlocal current
        if not current:
            return
        questions.append(
            CivilQuestion(
                number=int(current["number"]),
                section=str(current["section"]),
                statement=space_text(str(current["statement"])),
                answer=str(current.get("answer", "")).strip(),
                basis=space_text(str(current.get("basis", "")).strip()),
                explanation=space_text(str(current.get("explanation", "")).strip()),
                source=normalize_text(str(current.get("source", "")).strip()),
            )
        )
        current = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if not in_body:
            if "OX 본문" in text:
                in_body = True
            continue

        if text in SECTION_ORDER:
            flush()
            section = text
            continue

        if current is None and section and not re.match(r"^\d+\.\s*", text):
            # section intro memo
            continue

        match = re.match(r"^(\d+)\.\s*(.*)$", text)
        if match:
            flush()
            current = {
                "number": int(match.group(1)),
                "section": section,
                "statement": match.group(2).strip(),
            }
            continue

        if not current:
            continue
        if text.startswith("● 정답:") or text.startswith("정답:"):
            current["answer"] = text.split(":", 1)[1].strip()
        elif text.startswith("근거:"):
            current["basis"] = text.split(":", 1)[1].strip()
        elif text.startswith("해설:"):
            current["explanation"] = text.split(":", 1)[1].strip()
        elif text.startswith("출처:"):
            current["source"] = text.split(":", 1)[1].strip()

    flush()
    return questions


def extract_quoted_statements(text: str) -> list[str]:
    return [space_text(m) for m in re.findall(r'"([^"]+)"', text)]


def extract_basis_focus(basis: str) -> str:
    text = space_text(basis)
    match = re.search(r"\(([^()]+)\)", text)
    if match:
        return match.group(1).strip()
    match = re.search(r"(제\d+조(?:의\d+)?(?:\s*제\d+항)?)", text)
    if match:
        return match.group(1).replace("  ", " ").strip()
    return ""


def extract_statement_focus(statement: str) -> str:
    words: list[str] = []
    for token in KIWI.tokenize(space_text(statement)):
        if not token.tag.startswith("NN"):
            continue
        word = token.form.strip()
        if len(word) <= 1 or word in FOCUS_STOPWORDS or word in words:
            continue
        words.append(word)
        if len(words) == 3:
            break
    return "·".join(words)


def select_focus(basis: str, statement: str) -> str:
    basis_focus = extract_basis_focus(basis)
    if basis_focus and not re.fullmatch(r"제\d+조(?:의\d+)?(?:\s*제\d+항)?", basis_focus):
        return basis_focus
    statement_focus = extract_statement_focus(statement)
    return statement_focus or basis_focus


def build_focus_key(focus: str, fallback: str) -> str:
    if focus:
        return f"이 문제는 '{focus}' 기준 문언을 정확히 기억하는지가 핵심이다."
    return fallback


def build_focus_example(focus: str, fallback: str) -> str:
    if focus:
        return f"예를 들어 '{focus}' 문제는 조문 제목이나 조문 번호를 먼저 보고, 그다음 요건과 효과를 차례대로 떠올리면 정답이 훨씬 빨리 정리된다."
    return fallback


def build_explanation(q: CivilQuestion) -> str:
    text_for_match = " ".join([normalize_text(q.statement), normalize_text(q.basis)])
    guide = None
    for pattern, candidate in ISSUE_SPECIFIC_GUIDES:
        if pattern.search(text_for_match):
            guide = candidate
            break
    if guide is None:
        guide = SECTION_GUIDES.get(q.section, SECTION_GUIDES["기타"])
    focus = select_focus(q.basis, q.statement)
    key = build_focus_key(focus, guide["key"])
    raw = normalize_text(q.explanation)
    for pattern in GENERIC_EXPLANATION_PATTERNS:
        raw = re.sub(pattern, "", raw)
    raw = raw.strip()

    quotes = extract_quoted_statements(raw)
    if q.answer == "O":
        parts = [
            f"정답은 O입니다. 지문이 {q.basis or '민법·판례'} 기준과 일치한다.",
            key,
        ]
        return " ".join(part for part in parts if part).strip()
    else:
        manual = next((text for pattern, text in MANUAL_X_RULES if pattern.search(normalize_text(q.statement))), None)
        if manual:
            core = manual
        elif quotes:
            if len(quotes) == 1:
                core = f"정답은 X입니다. 올바른 정리는 “{quotes[0]}”이다."
            else:
                core = f"정답은 X입니다. 함께 기억할 올바른 정리는 “{quotes[0]}”, “{quotes[1]}”이다."
        else:
            basis = q.basis or "민법·판례"
            core = f"정답은 X입니다. 지문이 {basis} 기준의 요건·효과 또는 개념 연결과 맞지 않는다."

    return core.strip()


def set_default_font(doc: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(10.5 if style_name == "Normal" else 12)


def add_bullets(doc: Document, lines: list[str]) -> None:
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


def build_doc(questions: list[CivilQuestion], output_path: Path) -> None:
    doc = Document()
    set_default_font(doc)

    title = doc.add_heading("민법 OX 통합본", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("2014~2025 기출 반영 / 2026-04-16 기준 심화 보강판").bold = True

    intro = doc.add_paragraph()
    intro.add_run(
        "편집 기준: 기존 민법 OX 본문을 유지하면서 2025 기출 분석, 최근 출제형 변화, 2026 학습전략, "
        "박스형·케이스형 대비 포인트를 앞부분에 재구성했다. 본문은 현재 정제본을 바탕으로 다시 정렬해 "
        "단권화 자료로 바로 회독할 수 있게 정리했다."
    )

    doc.add_heading("Ⅰ. 민법 공부가 어려운 이유", level=1)
    add_bullets(
        doc,
        [
            "민법은 실생활에서 바로 쓰지 않는 법률용어가 많아 처음 진입 장벽이 높다. 법원, 제한능력자, 간주, 선의, 대항하지 못한다, 표현대리, 소급효, 담보책임, 채권자대위권 같은 용어에 익숙해지는 데 시간이 걸린다.",
            "민법은 독일식 총론·각론 체계를 계수한 과목이라, 뒤쪽 제도를 알아야 앞쪽 총론이 완전히 이해되는 구조를 가지고 있다. 그래서 1~2회독이 특히 힘들다.",
            "조문 양도 많고 범위도 넓다. 민법총칙 15문제, 채권총론 10문제, 채권각론 15문제가 최근 기준으로 출제되므로 총론만 하고 넘어가는 전략은 더 이상 통하지 않는다.",
        ],
    )

    doc.add_heading("Ⅱ. 2025 기출 분석", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "구분"
    hdr[1].text = "예전"
    hdr[2].text = "2024·2025"
    rows = [
        ("민법총칙", "12문제", "15문제"),
        ("채권법", "13문제", "25문제"),
        ("세부 구성", "-", "채권총론 10문제 / 채권각론 15문제"),
        ("총 문항 수", "25문제", "40문제"),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val

    add_bullets(
        doc,
        [
            "2025년 민법은 40문제로 확대된 뒤의 두 번째 시험으로, 2024년에 비해 비슷하거나 다소 쉬운 수준으로 평가되었다. 다만 25문제 시절보다 난도가 높아진 것은 분명하다.",
            "박스형 문제는 2025년에 9문제, 케이스형 문제는 6문제가 출제되어 단순 암기형보다 적용형·조합형 비중이 크게 늘었다.",
            "최근 민법은 중요한 조문과 판례를 정확히 알고 있는 수험생에게 유리하게 출제되고 있다. 결국 기본 조문과 판례를 확실히 반복한 수험생이 안정적으로 점수를 가져간다.",
        ],
    )

    doc.add_heading("Ⅲ. 2025 난이도 총평과 2026 대비 포인트", level=1)
    add_bullets(
        doc,
        [
            "2025년은 전체적으로 60점 이상은 맞출 수 있는 수준이라는 평가가 많았고, 무난하게 풀면 65~70점대도 가능한 시험으로 분석되었다.",
            "다만 채권법 비중 확대 이후에는 민법총칙만 강하게 잡고 채권을 얕게 보는 방식이 통하지 않는다. 2026년에는 채권총론·채권각론 비중을 더 무겁게 두는 편이 안전하다.",
            "문제 난도가 상대적으로 내려간 해에도 박스형과 케이스형이 시간을 잡아먹는다. ‘지문 자체가 어려운가’와 ‘시간 압박이 큰가’를 분리해서 연습해야 한다.",
            "예시로 표현대리, 보증채무, 임차인의 건물매수청구권 같은 쟁점은 판례 결론만 알아서는 부족하고, 요건과 효과를 정확히 적용할 수 있어야 한다.",
        ],
    )

    doc.add_heading("Ⅳ. 2026 학습요령", level=1)
    add_bullets(
        doc,
        [
            "민법은 과거보다 더 많은 시간을 배정해야 한다. 문제 수 확대와 채권법 비중 상승에 적응하려면 조기 착수가 유리하다.",
            "조문은 여전히 쉬운 점수를 주는 파트이므로 틀리면 손해가 크다. 조문집을 단권화 자료로 활용하면서 판례와 함께 반복하는 방식이 좋다.",
            "중요판례를 먼저 외우고, 시간이 허락하는 범위에서 나머지 판례를 확장하는 방식이 효율적이다. 최근 기출은 판례 결론 문구를 정확히 묻는 비중이 높다.",
            "기출문제는 조문별 반복 포인트를 보여주는 공개된 보물창고에 가깝다. 2014~2025 기출에서 어떤 논점이 반복되는지 먼저 파악한 뒤, 최신 판례와 미출제 중요 논점을 덧붙이는 방식이 가장 좋다.",
            "박스형 문제는 개별 지문 진위를 빠르게 판별하는 연습이 필요하고, 케이스형 문제는 사실관계와 당사자 지위를 먼저 표시하는 습관을 들여야 한다.",
        ],
    )

    doc.add_heading("Ⅴ. 비교해서 외우는 핵심 표", level=1)
    compare_table = doc.add_table(rows=1, cols=3)
    compare_table.style = "Table Grid"
    hdr = compare_table.rows[0].cells
    hdr[0].text = "비교 주제"
    hdr[1].text = "핵심 비교"
    hdr[2].text = "학습 포인트"
    for row_data in COMPARE_TABLE_ROWS:
        row = compare_table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value

    doc.add_heading("Ⅵ. 2014~2025 기출 함정 대비 암기 포인트", level=1)
    trap_table = doc.add_table(rows=1, cols=4)
    trap_table.style = "Table Grid"
    hdr = trap_table.rows[0].cells
    hdr[0].text = "혼동 포인트"
    hdr[1].text = "암기 내용"
    hdr[2].text = "두문자"
    hdr[3].text = "암기요령"
    for row_data in TRAP_MEMORY_ROWS:
        row = trap_table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value

    doc.add_heading("Ⅶ. OX 본문", level=1)
    grouped: dict[str, list[CivilQuestion]] = {section: [] for section in SECTION_ORDER}
    for q in questions:
        grouped.setdefault(q.section, []).append(q)

    for section in SECTION_ORDER:
        items = grouped.get(section, [])
        if not items:
            continue
        doc.add_paragraph(section, style="Heading 2")
        doc.add_paragraph(SECTION_NOTES[section])
        for q in items:
            doc.add_paragraph(f"{q.number}. {q.statement}")
            doc.add_paragraph(f"● 정답: {q.answer}")
            doc.add_paragraph(f"근거: {q.basis}")
            doc.add_paragraph(f"해설: {build_explanation(q)}")
            if q.source:
                doc.add_paragraph(f"출처: {q.source}")

    doc.add_heading("Ⅷ. 2026 기준 법령 확인 경로", level=1)
    add_bullets(
        doc,
        [
            "민법은 조문과 판례가 함께 움직이는 과목이므로, 시험 직전에는 최신 조문 개정 여부와 중요 판례 변동 여부를 함께 점검하는 것이 좋다.",
            "공식 조문 확인은 국가법령정보센터의 민법 본문에서 하는 것이 가장 안전하다.",
        ],
    )
    for url in LAW_URLS:
        doc.add_paragraph(url)

    doc.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx_path", nargs="?", default=str(DEFAULT_INPUT))
    args = parser.parse_args()

    output_path = Path(args.docx_path).expanduser()
    if not output_path.exists():
        raise SystemExit(f"입력 파일이 없습니다: {output_path}")

    backup = backup_file(output_path)
    questions = parse_existing_doc(output_path)
    build_doc(questions, output_path)
    print(f"backup={backup}")
    print(f"output={output_path}")
    print(f"questions={len(questions)}")


if __name__ == "__main__":
    main()
