#!/usr/bin/env python3
from __future__ import annotations

import argparse
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

from docx import Document


SECTION_TITLE = "문서심화분석 반영"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="문서 심화분석 결과를 DOCX/PDF에 자동 반영합니다.")
    parser.add_argument("--input-docx", type=Path, required=True, help="반영할 DOCX 문서")
    parser.add_argument("--analysis-file", type=Path, required=True, help="심화분석 결과 Markdown 파일")
    parser.add_argument("--output-pdf", type=Path, help="갱신할 PDF 경로. 기본값은 input-docx와 같은 이름의 PDF")
    parser.add_argument("--no-pdf", action="store_true", help="DOCX만 갱신하고 PDF는 만들지 않습니다.")
    parser.add_argument(
        "--section-title",
        default=SECTION_TITLE,
        help=f"문서 끝에 추가/교체할 반영 섹션 제목(기본: {SECTION_TITLE})",
    )
    return parser.parse_args()


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def backup_file(path: Path, label: str) -> Path:
    backup = path.with_name(f"{path.stem}_{label}_{timestamp()}{path.suffix}")
    shutil.copy2(path, backup)
    return backup


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def remove_existing_section(doc: Document, section_title: str) -> bool:
    start_idx = None
    for idx, para in enumerate(doc.paragraphs):
        if para.text.strip() == section_title:
            start_idx = idx
            break
    if start_idx is None:
        return False
    for para in list(doc.paragraphs[start_idx:]):
        delete_paragraph(para)
    return True


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        rows.append(cells)
    if len(rows) < 2:
        for line in lines:
            doc.add_paragraph(line)
        return
    header = rows[0]
    body = rows[2:] if len(rows) >= 2 and set(rows[1][0]) <= {"-", ":"} else rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for idx, cell in enumerate(header):
        table.rows[0].cells[idx].text = cell
    for row in body:
        tr = table.add_row().cells
        for idx, cell in enumerate(row[: len(header)]):
            tr[idx].text = cell


def apply_markdown_to_doc(doc: Document, markdown_text: str, section_title: str) -> None:
    remove_existing_section(doc, section_title)
    doc.add_page_break()
    doc.add_heading(section_title, level=1)

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            add_markdown_table(doc, block)
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        elif stripped[:2].isdigit() and stripped[2:4] in {". ", ") "}:
            doc.add_paragraph(stripped[3:].strip(), style="List Number")
        else:
            doc.add_paragraph(stripped)
        i += 1


def export_pdf_with_pages(input_docx: Path, output_pdf: Path) -> None:
    applescript = f"""
set inputPath to POSIX file "{input_docx.as_posix()}"
set pdfPath to POSIX file "{output_pdf.as_posix()}"

tell application "Pages"
\tactivate
\topen inputPath
\trepeat 30 times
\t\tif (count of documents) > 0 then exit repeat
\t\tdelay 1
\tend repeat
\tif (count of documents) is 0 then error "document did not open"
\tset docRef to front document
\texport docRef to pdfPath as PDF
\tclose docRef saving no
end tell
"""
    subprocess.run(["osascript", "-e", applescript], check=True)


def apply_analysis(input_docx: Path, analysis_file: Path, output_pdf: Path | None, *, section_title: str, make_pdf: bool) -> dict[str, str]:
    docx_path = input_docx.expanduser().resolve()
    analysis_path = analysis_file.expanduser().resolve()
    if not docx_path.exists():
        raise SystemExit(f"DOCX를 찾지 못했습니다: {docx_path}")
    if not analysis_path.exists():
        raise SystemExit(f"분석 파일을 찾지 못했습니다: {analysis_path}")

    backup_docx = backup_file(docx_path, "backup_before_apply_deep")
    doc = Document(docx_path)
    markdown_text = analysis_path.read_text(encoding="utf-8").strip()
    apply_markdown_to_doc(doc, markdown_text, section_title)
    doc.save(docx_path)

    result = {
        "docx": str(docx_path),
        "backup_docx": str(backup_docx),
    }
    if make_pdf:
        pdf_path = (output_pdf or docx_path.with_suffix(".pdf")).expanduser().resolve()
        backup_pdf = backup_file(pdf_path, "backup_before_apply_deep") if pdf_path.exists() else None
        export_pdf_with_pages(docx_path, pdf_path)
        result["pdf"] = str(pdf_path)
        if backup_pdf:
            result["backup_pdf"] = str(backup_pdf)
    return result


def main() -> int:
    args = parse_args()
    result = apply_analysis(
        args.input_docx,
        args.analysis_file,
        args.output_pdf,
        section_title=args.section_title,
        make_pdf=not args.no_pdf,
    )
    for key, value in result.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
