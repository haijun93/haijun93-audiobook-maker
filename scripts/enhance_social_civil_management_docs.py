#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from kiwipiepy import Kiwi


ROOT = Path(__file__).resolve().parent.parent
KIWI = Kiwi()

SUBJECT_CONFIG = {
    "social_insurance": {
        "input": Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/#Here/사회보험법_OX_통합본_2026_enhanced.docx"),
        "json": ROOT / ".work/social_insurance_ox_2026/social_insurance_ox_integrated_2026_final.json",
        "title": "사회보험법 OX 통합본",
        "subtitle": "2026 대비 심화 보강판",
    },
    "civil": {
        "input": Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/#Here/민법_OX_통합본_2026_cleaned_v1.docx"),
        "json": ROOT / ".work/civil_management_ox_2026/민법_OX_통합본_2026최종검수.json",
        "title": "민법 OX 통합본",
        "subtitle": "2014~2025 기출 반영 / 2026 대비 심화 보강판",
    },
    "management": {
        "input": Path("/Users/hyeokjunkong/Desktop/1차 시험/#STD/#Here/경영학_OX_통합본_2026_cleaned_v1.docx"),
        "json": ROOT / ".work/civil_management_ox_2026/경영학_OX_통합본_2026최종검수.json",
        "title": "경영학 OX 통합본",
        "subtitle": "2014~2025 기출 반영 / 2026 대비 심화 보강판",
    },
}

LOW_QUALITY_OX_STATEMENT_RE = re.compile(
    r"^(?:[ㄱ-ㅎ](?:\s*,\s*[ㄱ-ㅎ]){0,5}|[①-⑤](?:\s*,\s*[①-⑤]){0,5})$"
)

GENERIC_EXPLANATION_PATTERNS = (
    r"^정답은\s*[OX]입니다\.?$",
    r"^[OX]\.\s*.+?맞지 않는 진술입니다\.?$",
    r"^[OX]\.\s*.+?내용과 일치합니다\.?$",
    r"^현행 조문과 일치한다\.?$",
    r"^현행 시행령과 일치한다\.?$",
)

REGEX_REPLACEMENTS = (
    (r"노동\s*조합\s*및\s*노동\s*관계\s*조정법", "노동조합 및 노동관계조정법"),
    (r"근로자\s*참여\s*및\s*협력\s*증진에\s*관한\s*법률", "근로자참여 및 협력증진에 관한 법률"),
    (r"기간제\s*및\s*단시간\s*근로자\s*보호\s*등에\s*관한\s*법률", "기간제 및 단시간근로자 보호 등에 관한 법률"),
    (r"남녀\s*고용\s*평등과\s*일·\s*가정\s*양립\s*지원에\s*관한\s*법률", "남녀고용평등과 일·가정 양립 지원에 관한 법률"),
    (r"근로자\s*퇴직\s*급여\s*보장법", "근로자퇴직급여 보장법"),
    (r"파견\s*근로자\s*보호\s*등에\s*관한\s*법률", "파견근로자 보호 등에 관한 법률"),
    (r"외국인\s*근로자의\s*고용\s*등에\s*관한\s*법률", "외국인근로자의 고용 등에 관한 법률"),
    (r"최저\s*임금", "최저임금"),
    (r"적정\s*임금", "적정임금"),
    (r"근로\s*조건", "근로조건"),
    (r"근로\s*관계", "근로관계"),
    (r"직장\s*폐쇄", "직장폐쇄"),
    (r"해고\s*예고", "해고예고"),
    (r"근로자\s*공급\s*사업", "근로자공급사업"),
    (r"퇴직\s*급여", "퇴직급여"),
    (r"보장\s*법", "보장법"),
    (r"이행\s*강제금", "이행강제금"),
    (r"행정\s*소송", "행정소송"),
    (r"사회적\s*·\s*경제적", "사회적·경제적"),
    (r"사회\s*보장", "사회보장"),
    (r"지방\s*자치\s*단체", "지방자치단체"),
    (r"고용\s*보험", "고용보험"),
    (r"산업\s*재해\s*보상\s*보험", "산업재해보상보험"),
    (r"국민\s*연금", "국민연금"),
    (r"국민\s*건강\s*보험", "국민건강보험"),
    (r"보험\s*료", "보험료"),
    (r"수급\s*권", "수급권"),
    (r"피보험자\s*격", "피보험자격"),
    (r"재심\s*사", "재심사"),
    (r"상임\s*위원", "상임위원"),
    (r"직장\s*내", "직장 내"),
    (r"중앙\s*노동\s*위원회", "중앙노동위원회"),
    (r"지방\s*노동\s*위원회", "지방노동위원회"),
    (r"노동\s*위원회", "노동위원회"),
    (r"노사\s*협의회", "노사협의회"),
    (r"고충\s*처리\s*위원", "고충처리위원"),
    (r"최저\s*임금\s*위원회", "최저임금위원회"),
    (r"건강보험\s*분쟁\s*조정\s*위원회", "건강보험분쟁조정위원회"),
    (r"국민연금\s*기금\s*운용\s*위원회", "국민연금기금운용위원회"),
    (r"업무상\s*질병\s*판정\s*위원회", "업무상질병판정위원회"),
    (r"외국인력\s*정책\s*위원회", "외국인력정책위원회"),
    (r"근로\s*시간\s*면제\s*심의\s*위원회", "근로시간면제심의위원회"),
    (r"산업안전보건\s*위원회", "산업안전보건위원회"),
    (r"고용\s*노동부\s*장관", "고용노동부장관"),
    (r"사회보장\s*위원회", "사회보장위원회"),
    (r"제\s*(\d+)\s*조", r"제\1조"),
    (r"제\s*(\d+)\s*항", r"제\1항"),
    (r"제\s*(\d+)\s*호", r"제\1호"),
)

SOCIAL_COMMITTEE_NOTES = {
    "사회보장위원회": "사회보장위원회 포인트: 국무총리 소속이고, 사회보장 기본계획과 주요 사회보장정책을 심의·조정하는 컨트롤타워다. 보건복지부장관 소속으로 바꾸거나 실무위원회·전문위원회 존재를 부정하는 함정이 자주 나온다.",
    "고용보험위원회": "고용보험위원회 포인트: 고용보험 제도와 기금운용의 주요 사항을 심의하는 기구다. 위원 구성은 근로자·사용자·공익·정부 대표를 포괄하므로 단순 삼자구성으로 오해하지 않도록 주의해야 한다.",
    "업무상질병판정위원회": "업무상질병판정위원회 포인트: 고용노동부장관 소속이 아니라 근로복지공단 소속기관에 두는 것이 핵심이다. 시험에서는 소속 기관과 위원장 주체를 바꾸는 문제가 자주 나온다.",
    "국민연금기금운용위원회": "국민연금기금운용위원회 포인트: 국민연금 제도 일반을 심의하는 국민연금심의위원회와 달리, 기금 운용 사항을 심의·의결하는 별도 기구다. 위원장은 보건복지부장관이다.",
    "건강보험분쟁조정위원회": "건강보험분쟁조정위원회 포인트: 건강보험 자격·보험료·급여 관련 불복사건을 심리하는 재심 단계 기구다. 건강보험심사평가원이나 국민건강보험공단 내부 기구로 착각하는 함정에 주의해야 한다.",
    "건강보험정책심의위원회": "건강보험정책심의위원회 포인트: 건강보험 정책과 보험료율, 수가 구조 등 핵심 정책사항을 심의한다. 분쟁조정위원회와 기능을 뒤바꾸는 문제가 자주 출제된다.",
    "산업재해보상보험재심사위원회": "산업재해보상보험재심사위원회 포인트: 산재보험 재심 단계의 독립적 불복기구다. 공단 내부 심사와 재심사위원회를 혼동하는지 여부가 자주 출제된다.",
}

CIVIL_TOPIC_NOTES = {
    "총칙·권리능력·법원": "민법 총칙은 정의와 순서를 바꾸는 함정이 많다. 법원·관습법·권리능력 문제는 조문 순서와 판례 결론을 함께 기억하는 편이 좋다.",
    "의사표시·법률행위·대리": "민법은 요건과 효과를 미세하게 바꾸는 방식으로 오답을 만든다. 무효·취소·추인·대리권 남용을 한 묶음으로 비교해야 한다.",
    "물건·물권": "물권법은 공시·등기·점유·선의취득처럼 요건이 누적되는 구조가 많아 한 요소라도 빠지면 결론이 달라진다.",
    "채권·계약·불법행위": "최근 3년 비중이 가장 높은 영역이다. 계약유형별 효과, 동시이행·해제·손해배상, 불법행위 요건을 사례형으로 연결하는 문제가 많다.",
    "법인": "법인은 대표권 제한, 불법행위책임, 비법인사단과의 구별이 핵심이다.",
    "친족·상속": "친족·상속은 출제 비중은 크지 않지만 조문 문언을 정확히 묻는 문제가 많아 용어를 그대로 외우는 편이 효율적이다.",
    "기타": "민법 기타 파트는 조문과 판례가 섞여 있으므로, 같은 쟁점에서 ‘누가 어떤 효과를 주장할 수 있는가’를 중심으로 정리하는 것이 좋다.",
}

MANAGEMENT_TOPIC_NOTES = {
    "경영일반·전략": "전략 파트는 개념구분형과 이론가·모형형 함정이 자주 나온다. 원가우위·차별화·집중전략을 서로 바꾸는 문제가 대표적이다.",
    "조직행동·인사": "동기부여이론, 리더십, 권력, 조직문화처럼 이름이 비슷한 이론을 뒤섞는 문제가 많다. 핵심 개념과 대표 학자를 같이 묶어두는 것이 좋다.",
    "마케팅": "세분화·표적시장·포지셔닝, 제품수명주기, 소비자행동처럼 개념 간 경계를 흐리는 문제가 반복된다.",
    "생산·운영관리": "대기모형, 재고모형, 품질관리, 공정설계 등 계산·수치형 함정이 강하다. 공식과 해석을 같이 외워야 안정적이다.",
    "재무·회계": "재무·회계는 계산형과 개념형이 혼합된다. 지표의 의미와 증감 방향까지 함께 이해해야 오답을 줄일 수 있다.",
    "경영정보·데이터": "정보시스템 파트는 용어가 비슷해 보이지만 기능과 목적이 다르다. 약어와 정의를 세트로 기억하는 것이 효율적이다.",
    "기타": "경영학 기타 파트는 다양한 세부주제가 혼합되므로 문제 stem의 핵심 키워드를 먼저 잡는 연습이 중요하다.",
}


@dataclass
class QuestionBlock:
    number: int
    section: str
    statement: str
    answer: str
    basis: str
    explanation: str
    source: str


def normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = text.replace("•", "·").replace("ㆍ", "·")
    text = text.replace("“", '"').replace("”", '"').replace("’", "'").replace("‘", "'")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def compact_text(text: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", normalize_text(text))


def needs_spacing(text: str) -> bool:
    compact = re.sub(r"[^가-힣]", "", text)
    if len(compact) >= 12 and " " not in text:
        return True
    if re.search(r"[가-힣]{8,}", text):
        return True
    return False


def apply_replacements(text: str) -> str:
    updated = normalize_text(text)
    for pattern, replacement in REGEX_REPLACEMENTS:
        updated = re.sub(pattern, replacement, updated)
    updated = re.sub(r"\s*·\s*", "·", updated)
    updated = re.sub(r"\s{2,}", " ", updated)
    updated = re.sub(r"\(\s+", "(", updated)
    updated = re.sub(r"\s+\)", ")", updated)
    updated = re.sub(r"\s+([.,;:])", r"\1", updated)
    return updated.strip()


def space_text(text: str) -> str:
    updated = normalize_text(text)
    if needs_spacing(updated):
        updated = KIWI.space(updated)
    return apply_replacements(updated)


def is_low_quality_statement(statement: str) -> bool:
    normalized = normalize_text(statement)
    return bool(LOW_QUALITY_OX_STATEMENT_RE.fullmatch(normalized))


def has_sentence_predicate(statement: str) -> bool:
    return bool(
        re.search(
            r"(다|한다|된다|있다|없다|받는다|진다|필요하다|해당한다|가능하다|불가능하다|아니다|할 수 있다|할 수 없다)\.?$",
            normalize_text(statement),
        )
    )


def is_context_free_fragment(statement: str) -> bool:
    normalized = normalize_text(statement)
    compact = compact_text(normalized)
    if not normalized:
        return False
    if re.fullmatch(r"[0-9,]+(개|명|원|만원|일|개월|년)", compact):
        return True
    if normalized.startswith(("ㄱ", "①", "경제적 책임")) and len(compact) <= 40:
        return True
    if not has_sentence_predicate(normalized) and len(compact) <= 45:
        return True
    return False


def parse_answer_value(text: str) -> str:
    match = re.match(r"^(?:●\s*)?정답:\s*(.*)$", normalize_text(text))
    return match.group(1).strip() if match else ""


def is_bad_choice_payload(text: str) -> bool:
    normalized = normalize_text(text)
    compact = compact_text(normalized)
    if not normalized:
        return True
    if is_low_quality_statement(normalized) or is_context_free_fragment(normalized):
        return True
    if re.fullmatch(r"[ㄱ-ㅎ①-⑤,\-\s]+", normalized):
        return True
    if len(compact) <= 20 and not has_sentence_predicate(normalized):
        return True
    return False


def is_combo_choice_statement(text: str) -> bool:
    normalized = normalize_text(text)
    if re.match(r"^[ㄱ-ㅎ①-⑤](?:\s*,\s*[ㄱ-ㅎ①-⑤]){1,5}\s*은", normalized):
        return True
    if any(keyword in normalized for keyword in ("모두 고른 것은", "옳게 짝지어진 것은", "바르게 연결된 것은")):
        return True
    return False


def is_generic_explanation(text: str) -> bool:
    normalized = normalize_text(text)
    if not normalized:
        return True
    for pattern in GENERIC_EXPLANATION_PATTERNS:
        if re.match(pattern, normalized):
            return True
    return False


def backup_file(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}_backup_before_enhance_{timestamp}{path.suffix}")
    backup.write_bytes(path.read_bytes())
    return backup


def parse_doc(path: Path) -> tuple[list[str], list[QuestionBlock]]:
    doc = Document(str(path))
    preamble: list[str] = []
    items: list[QuestionBlock] = []
    section = ""
    current: dict[str, str | int] | None = None
    started_questions = False

    def flush() -> None:
        nonlocal current
        if current:
            items.append(
                QuestionBlock(
                    number=int(current["number"]),
                    section=str(current["section"]),
                    statement=str(current["statement"]),
                    answer=str(current.get("answer", "")),
                    basis=str(current.get("basis", "")),
                    explanation=str(current.get("explanation", "")),
                    source=str(current.get("source", "")),
                )
            )
            current = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if text.startswith("■ "):
            section = text[2:].strip()
            if not started_questions:
                preamble.append(text)
            continue
        match = re.match(r"^(\d+)\.\s*(.*)$", text)
        if match:
            started_questions = True
            flush()
            current = {
                "number": int(match.group(1)),
                "section": section,
                "statement": match.group(2),
            }
            continue
        if current is None:
            preamble.append(text)
            continue
        if text.startswith("● 정답:"):
            current["answer"] = text.split(":", 1)[1].strip()
        elif text.startswith("근거:"):
            current["basis"] = text.split(":", 1)[1].strip()
        elif text.startswith("해설:"):
            current["explanation"] = text.split(":", 1)[1].strip()
        elif text.startswith("출처:"):
            current["source"] = text.split(":", 1)[1].strip()
    flush()
    return preamble, items


def parse_social_doc(path: Path) -> tuple[list[str], list[QuestionBlock]]:
    doc = Document(str(path))
    preamble: list[str] = []
    items: list[QuestionBlock] = []
    section = ""
    current: dict[str, str | int] | None = None
    in_body = False

    def flush() -> None:
        nonlocal current
        if current:
            items.append(
                QuestionBlock(
                    number=int(current["number"]),
                    section=str(current["section"]),
                    statement=str(current["statement"]),
                    answer=str(current.get("answer", "")),
                    basis=str(current.get("basis", "")),
                    explanation=str(current.get("explanation", "")),
                    source=str(current.get("source", "")),
                )
            )
            current = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if not in_body:
            preamble.append(text)
            if text == "4. 통합 OX 본문":
                in_body = True
            continue
        if text.startswith("■ "):
            section = text[2:].strip()
            continue
        match = re.match(r"^(\d+)\.\s*(.*)$", text)
        if match:
            flush()
            current = {
                "number": int(match.group(1)),
                "section": section,
                "statement": match.group(2),
            }
            continue
        if current is None:
            continue
        answer = parse_answer_value(text)
        if answer:
            current["answer"] = answer
        elif text.startswith("근거:"):
            current["basis"] = text.split(":", 1)[1].strip()
        elif text.startswith("해설:"):
            current["explanation"] = text.split(":", 1)[1].strip()
        elif text.startswith("출처:"):
            current["source"] = text.split(":", 1)[1].strip()
    flush()
    return preamble, items


def extract_choice_payload(statement: str) -> str:
    match = re.search(r"\[선택지\s*[^:\]]+\s*:\s*(.+?)\]\s*$", normalize_text(statement))
    if match:
        return space_text(match.group(1))
    return space_text(statement)


def extract_stem(statement: str) -> str:
    normalized = normalize_text(statement)
    return normalize_text(re.sub(r"\[선택지\s*[^:\]]+\s*:\s*.+?\]\s*$", "", normalized))


def stem_is_negative(stem: str) -> bool:
    negative_patterns = (
        r"옳지\s*않은\s*것",
        r"옳지않은것",
        r"해당하지\s*않은\s*것",
        r"해당하지않은것",
        r"아닌\s*것",
        r"아닌것",
        r"틀린\s*것",
        r"부적절한\s*것",
    )
    return any(re.search(pattern, stem) for pattern in negative_patterns)


def build_true_choice_map(entries: list[dict[str, object]]) -> dict[str, list[str]]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for entry in entries:
        grouped[str(entry["source"])].append(entry)
    result: dict[str, list[str]] = {}
    for source, group in grouped.items():
        stem = extract_stem(str(group[0]["statement"]))
        negative = stem_is_negative(stem)
        if negative:
            truth_entries = [entry for entry in group if str(entry["answer"]).upper() == "X"]
        else:
            truth_entries = [entry for entry in group if str(entry["answer"]).upper() == "O"]
        truths = [extract_choice_payload(str(entry["statement"])) for entry in truth_entries]
        if truths:
            result[source] = truths
    return result


def compute_recent_trends(entries: list[dict[str, object]], subject: str) -> tuple[list[tuple[str, Counter]], dict[str, dict[str, int]]]:
    recent = [entry for entry in entries if str(entry.get("year", "")) in {"2023", "2024", "2025"}]
    year_topic: list[tuple[str, Counter]] = []
    for year in ("2023", "2024", "2025"):
        year_topic.append((year, Counter(str(entry["topic"]) for entry in recent if str(entry["year"]) == year)))

    if subject == "civil":
        patterns = {
            "판례형": re.compile(r"(판례|다툼이 있으면 판례|대법원)"),
            "요건·효과형": re.compile(r"(요한다|필요하다|효력이|취소|무효|해제|해지|추인|철회)"),
            "주체·상대방형": re.compile(r"(채권자|채무자|제3자|당사자|대리인|본인|상속인|점유자)"),
            "기간·수치형": re.compile(r"(\d+년|\d+개월|\d+일|소멸시효|제척기간)"),
            "사례형": re.compile(r"(甲|乙|병|정)"),
        }
    else:
        patterns = {
            "개념구분형": re.compile(r"(해당한다|해당하지|유형|특성|원천|구성요소|차원)"),
            "이론가·모형형": re.compile(r"(Porter|Maslow|Mintzberg|Herzberg|Ansoff|BSC|SWOT|가치사슬|균형성과표)"),
            "계산·수치형": re.compile(r"(\d+%|\d+개|\d+명|\d+년|\d+원|\d+배)"),
            "사례판단형": re.compile(r"(갑회사|A기업|B기업|사례|기업이)"),
            "비교형": re.compile(r"(차이|비교|구분|장점|단점|아닌 것)"),
        }
    trap_counts: dict[str, dict[str, int]] = {}
    for year in ("2023", "2024", "2025"):
        items = [entry for entry in recent if str(entry["year"]) == year]
        trap_counts[year] = {name: sum(1 for entry in items if pattern.search(str(entry["statement"]))) for name, pattern in patterns.items()}
    return year_topic, trap_counts


def classify_trap(statement: str, subject: str) -> str:
    text = normalize_text(statement)
    if subject == "civil":
        mapping = [
            ("판례형", r"(판례|다툼이 있으면 판례|대법원)"),
            ("요건·효과형", r"(요한다|필요하다|효력이|취소|무효|해제|해지|추인|철회)"),
            ("주체·상대방형", r"(채권자|채무자|제3자|당사자|대리인|본인|상속인|점유자)"),
            ("기간·수치형", r"(\d+년|\d+개월|\d+일|소멸시효|제척기간)"),
            ("사례형", r"(甲|乙|병|정)"),
        ]
    elif subject == "management":
        mapping = [
            ("이론가·모형형", r"(Porter|Maslow|Mintzberg|Herzberg|Ansoff|BSC|SWOT|가치사슬|균형성과표)"),
            ("계산·수치형", r"(\d+%|\d+개|\d+명|\d+년|\d+원|\d+배)"),
            ("사례판단형", r"(갑회사|A기업|B기업|사례|기업이)"),
            ("비교형", r"(차이|비교|구분|장점|단점|아닌 것)"),
            ("개념구분형", r"(해당한다|해당하지|유형|특성|원천|구성요소|차원)"),
        ]
    else:
        mapping = [
            ("위원회·절차형", r"(위원회|심의|심사|재심|조정|위원)"),
            ("기간·수치형", r"(\d+일|\d+개월|\d+년|과반수|3분의 2|100분의|이상|이내|초과)"),
            ("주체·권한형", r"(장관|공단|공단|근로자|사용자|수급권자|위원회)"),
            ("자격·급여형", r"(수급권|피보험자|가입자|유족연금|보험료|급여)"),
        ]
    for label, pattern in mapping:
        if re.search(pattern, text):
            return label
    return "문언형"


def trap_comment(label: str, subject: str) -> str:
    if subject == "civil":
        return {
            "판례형": "최근 기출에서는 판례 결론 문구를 조금만 비틀어도 오답이 되는 문제가 많으므로 결론과 이유를 함께 기억해야 한다.",
            "요건·효과형": "민법은 요건 하나가 빠지면 효과가 완전히 달라지므로 요건과 법률효과를 세트로 정리해야 한다.",
            "주체·상대방형": "당사자·제3자·대리인·상속인의 지위를 뒤바꾸는 함정이 많아, 누가 무엇을 주장할 수 있는지 먼저 확인하는 습관이 중요하다.",
            "기간·수치형": "소멸시효·제척기간·기간 계산은 숫자와 기산점을 함께 외워야 안정적으로 맞힌다.",
            "사례형": "사례형 문제는 사실관계 한 줄 차이로 결론이 달라지므로 당사자 지위와 시간 순서를 먼저 잡는 것이 좋다.",
            "문언형": "민법은 조문 문언을 한두 단어 바꾸는 기본형 함정도 꾸준히 출제된다.",
        }[label]
    if subject == "management":
        return {
            "이론가·모형형": "최근 기출에서는 학자 이름과 모형의 핵심 요소를 교차시켜 오답을 만드는 방식이 자주 보인다.",
            "계산·수치형": "계산형은 공식을 아는 것만으로 부족하고, 숫자의 의미와 해석 방향까지 함께 정리해야 한다.",
            "사례판단형": "사례판단형은 개념 정의를 실제 기업 상황에 적용할 수 있는지를 묻는 경우가 많다.",
            "비교형": "유사 개념의 차이와 장단점을 한 문제 안에서 섞는 유형이 반복되므로 비교표 정리가 효과적이다.",
            "개념구분형": "핵심 개념의 경계를 흐리게 하는 기본형 함정이 많으니 정의를 짧게라도 정확히 외우는 것이 좋다.",
            "문언형": "용어 하나만 바꾼 기본형 OX도 꾸준히 출제되므로 핵심 키워드를 정확히 기억해야 한다.",
        }[label]
    return {
        "위원회·절차형": "최근 사회보험법 기출은 위원회 소속, 단계별 불복절차, 의결구조를 바꾸는 방식의 함정이 특히 강하다.",
        "기간·수치형": "보험료·급여·신청기간은 숫자와 기산점이 함께 출제되므로 하나만 외우면 오답이 되기 쉽다.",
        "주체·권한형": "장관, 공단, 위원회, 가입자 등 주체를 바꾸는 함정이 매우 많으므로 권한 주체를 먼저 확인해야 한다.",
        "자격·급여형": "자격요건과 급여요건을 뒤섞는 문제가 자주 나오므로 대상·요건·효과를 묶어 정리하는 편이 좋다.",
        "문언형": "조문 문언을 한두 단어만 바꿔도 결론이 달라지는 기본형 함정이 반복된다.",
    }[label]


def choose_topic_note(topic: str, subject: str) -> str:
    if subject == "civil":
        return CIVIL_TOPIC_NOTES.get(topic, CIVIL_TOPIC_NOTES["기타"])
    if subject == "management":
        return MANAGEMENT_TOPIC_NOTES.get(topic, MANAGEMENT_TOPIC_NOTES["기타"])
    return ""


def build_choice_explanation(item: QuestionBlock, truth_map: dict[str, list[str]], subject: str) -> str:
    trap = classify_trap(item.statement, subject)
    topic_note = choose_topic_note(item.section, subject)
    truths = truth_map.get(item.source, [])
    normalized_statement = compact_text(item.statement)
    other_truths = [
        text
        for text in truths
        if compact_text(text) != normalized_statement and not is_bad_choice_payload(text)
    ]
    if item.answer == "O":
        parts = [
            "정답은 O입니다.",
            "같은 기출문항에서 옳은 진술로 기능하는 문장을 그대로 OX화한 것이다.",
            topic_note,
            trap_comment(trap, subject),
        ]
        return " ".join(part for part in parts if part)
    if other_truths:
        if len(other_truths) == 1:
            lead = f'정답은 X입니다. 같은 기출문항에서 옳은 진술로는 "{other_truths[0]}"가 제시되었다.'
        else:
            lead = "정답은 X입니다. 같은 기출문항에서 옳은 진술로는 " + ", ".join(f'"{text}"' for text in other_truths[:2]) + " 등이 제시되었다."
    else:
        lead = "정답은 X입니다. 이 진술은 같은 쟁점의 기준 문언·판례 결론과 맞지 않는다."
    return " ".join(part for part in [lead, topic_note, trap_comment(trap, subject)] if part)


def build_social_explanation(item: QuestionBlock, entry_map: dict[int, dict[str, object]]) -> str:
    trap = classify_trap(item.statement, "social_insurance")
    entry = entry_map.get(item.number)
    base_note = normalize_text(str(entry.get("note", ""))) if entry else normalize_text(item.explanation)
    if not base_note or is_generic_explanation(base_note):
        if item.answer == "O":
            base_note = f"정답은 O입니다. {item.basis}의 현행 기준과 일치하는 진술이다."
        else:
            base_note = f"정답은 X입니다. {item.basis}의 현행 기준과 어긋나는 부분을 조정한 문항이다."
    committee_note = ""
    for key, note in SOCIAL_COMMITTEE_NOTES.items():
        if key in item.statement:
            committee_note = note
            break
    return " ".join(part for part in [space_text(base_note), trap_comment(trap, "social_insurance"), committee_note] if part)


def set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            s = doc.styles[style_name]
            s.font.name = "Malgun Gothic"
            s._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def build_choice_subject_doc(subject: str, input_path: Path, json_path: Path) -> tuple[int, Path]:
    backup = backup_file(input_path)
    _, items = parse_doc(input_path)
    entries = json.loads(json_path.read_text())["entries"]
    truth_map = build_true_choice_map(entries)
    year_topic, trap_counts = compute_recent_trends(entries, subject)

    cleaned: list[QuestionBlock] = []
    removed = 0
    for item in items:
        statement = space_text(item.statement)
        payload = extract_choice_payload(statement)
        if (
            is_low_quality_statement(statement)
            or is_context_free_fragment(statement)
            or is_bad_choice_payload(payload)
            or is_combo_choice_statement(statement)
        ):
            removed += 1
            continue
        cleaned.append(
            QuestionBlock(
                number=len(cleaned) + 1,
                section=space_text(item.section),
                statement=statement,
                answer=item.answer.strip(),
                basis=space_text(item.basis),
                explanation=item.explanation,
                source=normalize_text(item.source),
            )
        )

    doc = Document()
    set_default_font(doc)
    title = doc.add_heading(SUBJECT_CONFIG[subject]["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(SUBJECT_CONFIG[subject]["subtitle"]).bold = True
    intro = doc.add_paragraph()
    intro.add_run(
        f"편집 기준: 최근 3년 출제경향과 함정 유형을 반영해, 단독 학습에 도움이 되지 않는 저품질 문항을 정리하고 해설을 '왜 맞고 왜 틀리는지' 중심으로 보강했다. "
        f"이번 업데이트에서는 저품질·문맥절단 문항 {removed}개를 삭제했다."
    )

    doc.add_heading("Ⅰ. 최근 3년 출제 경향", level=1)
    for year, counter in year_topic:
        top = ", ".join(f"{topic} {count}개" for topic, count in counter.most_common(6))
        doc.add_paragraph(f"{year}: {top}")

    if subject == "civil":
        bullets = [
            "민법은 최근 3년 모두 채권·계약·불법행위 비중이 가장 높았고, 의사표시·법률행위·대리 파트가 그 뒤를 이었다.",
            "판례형과 사례형 출제가 압도적으로 많아, 조문 암기만으로는 부족하고 판례 결론 문구를 정확히 익혀야 한다.",
            "특히 주체·상대방을 바꾸는 함정과 요건·효과를 뒤집는 함정이 반복되므로, 누구에게 어떤 효과가 귀속되는지를 먼저 확인하는 습관이 중요하다.",
        ]
    else:
        bullets = [
            "경영학은 개념구분형이 가장 많고, 2025년에는 계산·수치형과 이론가·모형형 비중이 함께 높아졌다.",
            "단순 정의형처럼 보여도 실제로는 유사 개념을 교차 배치해 오답을 유도하는 문제가 많다.",
            "특히 전략·조직행동·생산운영·재무회계는 '핵심 개념 + 대표 모형 + 숫자 해석'을 묶어 정리해야 효율적이다.",
        ]
    for bullet in bullets:
        doc.add_paragraph(bullet, style="List Bullet")

    trap_table = doc.add_table(rows=1, cols=4)
    trap_table.style = "Table Grid"
    hdr = trap_table.rows[0].cells
    hdr[0].text = "함정 유형"
    hdr[1].text = "2023"
    hdr[2].text = "2024"
    hdr[3].text = "2025"
    for trap_name in next(iter(trap_counts.values())).keys():
        row = trap_table.add_row().cells
        row[0].text = trap_name
        row[1].text = str(trap_counts["2023"][trap_name])
        row[2].text = str(trap_counts["2024"][trap_name])
        row[3].text = str(trap_counts["2025"][trap_name])

    doc.add_heading("Ⅱ. OX 본문", level=1)
    by_section: dict[str, list[QuestionBlock]] = defaultdict(list)
    for item in cleaned:
        by_section[item.section].append(item)

    for section, section_items in by_section.items():
        doc.add_heading(section, level=2)
        doc.add_paragraph(choose_topic_note(section, subject))
        for item in section_items:
            doc.add_paragraph(f"{item.number}. {item.statement}")
            doc.add_paragraph(f"● 정답: {item.answer}")
            doc.add_paragraph(f"근거: {item.basis}")
            doc.add_paragraph(f"해설: {build_choice_explanation(item, truth_map, subject)}")
            if item.source:
                doc.add_paragraph(f"출처: {item.source}")

    doc.save(str(input_path))
    return removed, backup


def build_social_insurance_doc(input_path: Path, json_path: Path) -> tuple[int, Path]:
    backup = backup_file(input_path)
    source = Document(str(input_path))
    obj = json.loads(json_path.read_text())
    entries = obj["entries"]
    entry_map = {int(entry["number"]): entry for entry in entries}

    doc = Document()
    set_default_font(doc)
    in_body = False
    current_number: int | None = None
    current_statement = ""
    current_basis = ""
    improved = 0
    for para in source.paragraphs:
        raw = para.text.strip()
        if not raw:
            continue
        normalized = apply_replacements(normalize_text(raw))
        if not in_body:
            doc.add_paragraph(normalized)
            if normalized == "4. 통합 OX 본문":
                in_body = True
            continue

        match = re.match(r"^(\d+)\.\s*(.*)$", raw)
        if match:
            current_number = int(match.group(1))
            current_statement = match.group(2).strip()
            current_basis = ""
            doc.add_paragraph(f"{current_number}. {space_text(current_statement)}")
            continue

        answer = parse_answer_value(raw)
        if answer or raw.startswith("정답:") or raw.startswith("● 정답:"):
            if current_number is not None:
                answer = str(entry_map.get(current_number, {}).get("answer", answer or "")).strip()
            doc.add_paragraph(f"정답: {answer}")
            continue

        if raw.startswith("근거:"):
            current_basis = raw.split(":", 1)[1].strip()
            doc.add_paragraph(f"근거: {apply_replacements(normalize_text(current_basis))}")
            continue

        if raw.startswith("해설:"):
            if current_number is None:
                doc.add_paragraph(f"해설: {apply_replacements(normalize_text(raw.split(':', 1)[1].strip()))}")
                continue
            item = QuestionBlock(
                number=current_number,
                section="",
                statement=current_statement,
                answer=str(entry_map.get(current_number, {}).get("answer", "")),
                basis=current_basis,
                explanation=raw.split(":", 1)[1].strip(),
                source=str(entry_map.get(current_number, {}).get("source", "")),
            )
            explanation = build_social_explanation(item, entry_map)
            if normalize_text(explanation) != normalize_text(item.explanation):
                improved += 1
            doc.add_paragraph(f"해설: {explanation}")
            continue

        if raw.startswith("출처:"):
            doc.add_paragraph(f"출처: {normalize_text(raw.split(':', 1)[1].strip())}")
            continue

        doc.add_paragraph(normalized)

    doc.save(str(input_path))
    return improved, backup


def main() -> None:
    parser = argparse.ArgumentParser(description="사회보험법·민법·경영학 OX 문서를 심화 보강합니다.")
    parser.add_argument(
        "subjects",
        nargs="*",
        choices=sorted(SUBJECT_CONFIG.keys()),
        default=["social_insurance", "civil", "management"],
    )
    args = parser.parse_args()

    for subject in args.subjects:
        config = SUBJECT_CONFIG[subject]
        if subject == "social_insurance":
            improved, backup = build_social_insurance_doc(config["input"], config["json"])
            print(f"{subject}: backup={backup} improved_explanations={improved}")
        else:
            removed, backup = build_choice_subject_doc(subject, config["input"], config["json"])
            print(f"{subject}: backup={backup} removed_low_quality={removed}")


if __name__ == "__main__":
    main()
