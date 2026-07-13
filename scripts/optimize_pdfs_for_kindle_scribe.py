#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DEFAULT_DOCX_FILES = [
    Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416/civil_law_OX_integrated_2026_v1.docx"),
    Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416/labor_law_OX_integrated_2026_v1.docx"),
    Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416/management_OX_integrated_2026_v1.docx"),
    Path(str(Path.home()) + "/Desktop/1차 시험/#STD/260416/social_insurance_law_OX_integrated_2026_v1.docx"),
]

# 6-inch Kindle profile.
# Smaller fixed-layout pages reduce right-edge clipping on non-Scribe devices.
PROFILE_SETTINGS = {
    "scribe": {
        "page_width": Inches(6.2),
        "page_height": Inches(8.27),
        "margin_lr": Inches(0.32),
        "margin_tb": Inches(0.38),
        "title_space_before": Pt(220),
        "title_space_after": Pt(18),
        "subtitle_text": "Kindle Scribe Reading Edition",
        "subtitle_size": Pt(10.5),
        "normal_size": 11.5,
        "title_size": 16,
        "h1_size": 13.5,
        "h2_size": 12.5,
        "h3_size": 11.5,
        "line_spacing": 1.15,
        "para_space_after": Pt(2),
        "title_space_after_para": Pt(6),
        "heading_space_before": Pt(5),
        "heading_space_after": Pt(2),
        "answer_indent": Inches(0.04),
        "numbered_space_before": Pt(2),
        "table_line_spacing": 1.05,
        "table_min_font": Pt(10.5),
    },
    "kindle6": {
        "page_width": Inches(3.58),
        "page_height": Inches(4.84),
        "margin_lr": Inches(0.18),
        "margin_tb": Inches(0.20),
        "title_space_before": Pt(120),
        "title_space_after": Pt(10),
        "subtitle_text": "Kindle 6-inch Reading Edition",
        "subtitle_size": Pt(8.5),
        "normal_size": 9.4,
        "title_size": 13,
        "h1_size": 10.8,
        "h2_size": 10.0,
        "h3_size": 9.4,
        "line_spacing": 1.05,
        "para_space_after": Pt(1),
        "title_space_after_para": Pt(4),
        "heading_space_before": Pt(3),
        "heading_space_after": Pt(1),
        "answer_indent": Inches(0.02),
        "numbered_space_before": Pt(1),
        "table_line_spacing": 1.0,
        "table_min_font": Pt(8.4),
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="DOCX를 Kindle 프로필에 맞는 읽기용 레이아웃으로 재구성한 뒤 PDF로 다시 내보냅니다."
    )
    parser.add_argument("docx_files", nargs="*", type=Path, help="최적화할 DOCX 파일들")
    parser.add_argument(
        "--profile",
        choices=sorted(PROFILE_SETTINGS),
        default="scribe",
        help="출력 PDF 대상 기기 프로필 (기본: scribe)",
    )
    return parser.parse_args()


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def set_style_font(doc: Document, style_name: str, *, size: float, bold: bool | None = None) -> None:
    if style_name not in doc.styles:
        return
    style = doc.styles[style_name]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def infer_subject_title(doc: Document, input_docx: Path) -> str:
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        match = re.match(r"^(.+?)\s+OX\s+통합본", text)
        if match:
            return match.group(1).strip()
        if text in {"사회보험법", "사회보장법", "노동법", "민법", "경영학", "경영학개론"}:
            return text

    stem = input_docx.stem
    stem = re.sub(r"[_-]+", " ", stem)
    stem = re.sub(r"\b\d{6,8}\b", "", stem)
    stem = re.sub(r"\b(v|ver)\s*\d+\b", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"\s+", " ", stem).strip()
    return stem or "학습자료"


def prepend_subject_title_page(doc: Document, subject_title: str, settings: dict) -> None:
    if not doc.paragraphs:
        anchor = doc.add_paragraph("")
    else:
        anchor = doc.paragraphs[0]

    page_break_para = anchor.insert_paragraph_before("")
    page_break_para.add_run().add_break(WD_BREAK.PAGE)

    title_para = page_break_para.insert_paragraph_before(subject_title, style="Title")
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = settings["title_space_before"]
    title_para.paragraph_format.space_after = settings["title_space_after"]

    subtitle_para = page_break_para.insert_paragraph_before(settings["subtitle_text"])
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_para.paragraph_format.space_after = Pt(0)
    for run in subtitle_para.runs:
        run.font.name = "Malgun Gothic"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        run.font.size = settings["subtitle_size"]


def optimize_docx_for_kindle(input_docx: Path, output_docx: Path, profile: str = "scribe") -> None:
    settings = PROFILE_SETTINGS[profile]
    doc = Document(str(input_docx))
    subject_title = infer_subject_title(doc, input_docx)
    prepend_subject_title_page(doc, subject_title, settings)

    for section in doc.sections:
        section.page_width = settings["page_width"]
        section.page_height = settings["page_height"]
        section.left_margin = settings["margin_lr"]
        section.right_margin = settings["margin_lr"]
        section.top_margin = settings["margin_tb"]
        section.bottom_margin = settings["margin_tb"]
        section.header_distance = Inches(0.2)
        section.footer_distance = Inches(0.2)

    set_style_font(doc, "Normal", size=settings["normal_size"])
    set_style_font(doc, "Title", size=settings["title_size"], bold=True)
    set_style_font(doc, "Heading 1", size=settings["h1_size"], bold=True)
    set_style_font(doc, "Heading 2", size=settings["h2_size"], bold=True)
    set_style_font(doc, "Heading 3", size=settings["h3_size"], bold=True)

    for para in iter_all_paragraphs(doc):
        pf = para.paragraph_format
        pf.line_spacing = settings["line_spacing"]
        pf.space_after = settings["para_space_after"]
        pf.space_before = Pt(0)

        text = para.text.strip()
        style_name = para.style.name if para.style is not None else ""
        if style_name == "Title":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = settings["title_space_after_para"]
        elif style_name in {"Heading 1", "Heading 2", "Heading 3"}:
            pf.keep_with_next = True
            pf.space_before = settings["heading_space_before"]
            pf.space_after = settings["heading_space_after"]
        elif text.startswith(("정답:", "● 정답:", "근거:", "해설:", "출처:")):
            pf.left_indent = settings["answer_indent"]
        elif text.startswith(tuple(f"{i}." for i in range(1, 10))) or text[:2].isdigit():
            pf.space_before = settings["numbered_space_before"]

        for run in para.runs:
            run.font.name = "Malgun Gothic"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    for table in doc.tables:
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    pf = para.paragraph_format
                    pf.line_spacing = settings["table_line_spacing"]
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    for run in para.runs:
                        run.font.name = "Malgun Gothic"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
                        if run.font.size is None or run.font.size.pt < settings["table_min_font"].pt:
                            run.font.size = settings["table_min_font"]

    doc.save(str(output_docx))


def export_pdf_with_pages(input_docx: Path, output_pdf: Path) -> None:
    applescript = f"""
set inputPath to POSIX file "{input_docx.as_posix()}"
set pdfPath to POSIX file "{output_pdf.as_posix()}"

tell application "Pages"
\tactivate
\topen inputPath
\trepeat 30 times
\t\tif (count of documents) > 0 then exit repeat
\t\tdelay 1
\tend repeat
\tif (count of documents) is 0 then error "document did not open"
\tset docRef to front document
\texport docRef to pdfPath as PDF
\tclose docRef saving no
end tell
"""
    subprocess.run(["osascript", "-e", applescript], check=True)


def backup_pdf(pdf_path: Path) -> Path | None:
    if not pdf_path.exists():
        return None
    backup = pdf_path.with_name(f"{pdf_path.stem}_backup_before_kindle_scribe_{timestamp()}{pdf_path.suffix}")
    backup.write_bytes(pdf_path.read_bytes())
    return backup


def inspect_pdf(pdf_path: Path) -> tuple[int, float, float]:
    with fitz.open(pdf_path) as pdf:
        page_count = pdf.page_count
        page = pdf.load_page(0)
        rect = page.rect
        return page_count, rect.width / 72.0, rect.height / 72.0


def optimize_one(docx_path: Path, profile: str) -> None:
    pdf_path = docx_path.with_suffix(".pdf")
    backup = backup_pdf(pdf_path)
    with tempfile.TemporaryDirectory() as tmpdir:
        temp_docx = Path(tmpdir) / docx_path.name
        optimize_docx_for_kindle(docx_path, temp_docx, profile=profile)
        export_pdf_with_pages(temp_docx, pdf_path)

    pages, width_in, height_in = inspect_pdf(pdf_path)
    print(f"docx={docx_path}")
    print(f"pdf={pdf_path}")
    print(f"backup={backup}")
    print(f"pages={pages}")
    print(f"page_size_in={width_in:.2f}x{height_in:.2f}")


def main() -> int:
    args = parse_args()
    docx_files = args.docx_files or DEFAULT_DOCX_FILES
    resolved = [Path(p).expanduser().resolve() for p in docx_files]

    for path in resolved:
        if not path.exists():
            raise SystemExit(f"입력 DOCX를 찾지 못했습니다: {path}")

    for path in resolved:
        optimize_one(path, args.profile)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
