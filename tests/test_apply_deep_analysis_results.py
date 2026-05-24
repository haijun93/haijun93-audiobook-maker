from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.apply_deep_analysis_results import apply_markdown_to_doc


class ApplyDeepAnalysisResultsTests(unittest.TestCase):
    def test_apply_markdown_replaces_existing_section(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.docx"
            doc = Document()
            doc.add_heading("원래 문서", level=1)
            doc.add_paragraph("본문")
            doc.add_heading("문서심화분석 반영", level=1)
            doc.add_paragraph("예전 분석")
            doc.save(path)

            reopened = Document(path)
            apply_markdown_to_doc(reopened, "# 새 분석\n\n- 핵심1\n- 핵심2", "문서심화분석 반영")
            reopened.save(path)

            final_doc = Document(path)
            texts = [p.text for p in final_doc.paragraphs if p.text.strip()]
            self.assertEqual(texts.count("문서심화분석 반영"), 1)
            self.assertIn("새 분석", texts)
            self.assertIn("핵심1", texts)
            self.assertIn("핵심2", texts)
            self.assertNotIn("예전 분석", texts)


if __name__ == "__main__":
    unittest.main()
